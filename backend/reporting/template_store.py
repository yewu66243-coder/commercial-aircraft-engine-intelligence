"""Local DOCX format templates, kept separate from the research document library."""
from __future__ import annotations

from collections import Counter
from copy import deepcopy
from datetime import datetime
from hashlib import sha256
from io import BytesIO
import json
from pathlib import Path
import re
from zipfile import BadZipFile, ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .format_profile import LATIN_REPORT_FONTS, normalize_chapters, normalize_report_format_profile

TEMPLATE_DIR = Path(__file__).resolve().parents[2] / "local_docs" / "report_templates"
MAX_TEMPLATE_BYTES = 10 * 1024 * 1024
BODY_MARKER = "{{REPORT_BODY}}"
STYLE_ROLES = {
    "body": "Normal", "title": "Title", "heading1": "Heading 1",
    "heading2": "Heading 2", "heading3": "Heading 3",
    "abstract": "Report Abstract", "keywords": "Report Keywords",
    "caption": "Report Caption", "source": "Report Source",
    "reference": "Report Reference", "contents": "Report Contents",
    "table": "Report Table", "list": "Report List",
}
EDITABLE_PATHS = {
    "font.body", "font.heading", "size.body", "size.heading1",
    "layout.margin_mm", "layout.line_spacing", "layout.cover", "layout.include_toc",
}


class ReportTemplateError(ValueError):
    pass


def template_path(template_id: str) -> Path:
    if not re.fullmatch(r"[a-f0-9]{24}", str(template_id)):
        raise ReportTemplateError("报告模板标识无效，请重新选择模板。")
    path = TEMPLATE_DIR / f"{template_id}.docx"
    if not path.is_file():
        raise ReportTemplateError("报告模板已不存在，请重新导入或选择其他模板。")
    return path


def get_template(template_id: str) -> dict:
    path = template_path(template_id).with_suffix(".json")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise ReportTemplateError("模板识别记录损坏，请重新导入模板。") from exc


def list_templates() -> list[dict]:
    items = []
    if TEMPLATE_DIR.is_dir():
        for path in TEMPLATE_DIR.glob("*.json"):
            try:
                item = get_template(path.stem)
                items.append(item)
            except ReportTemplateError:
                continue
    return sorted(items, key=lambda item: item["created_at"], reverse=True)


def _validate_package(data: bytes):
    if len(data) > MAX_TEMPLATE_BYTES:
        raise ReportTemplateError("Word 模板不能超过 10 MiB。")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(m.file_size for m in members) > 80 * 1024 * 1024:
                raise ReportTemplateError("模板解压后体积过大，请精简图片和正文后重新导入。")
            names = set(archive.namelist())
            if not {"word/document.xml", "[Content_Types].xml"} <= names:
                raise ReportTemplateError("文件不是有效的 DOCX 文档。")
            if any("vbaProject" in n or n.startswith("word/embeddings/") for n in names):
                raise ReportTemplateError("请使用不含宏或嵌入附件的 DOCX 模板。")
            # Header images are local relationships; external image/OLE references
            # must not be followed during automated document conversion.
            from lxml import etree
            parser = etree.XMLParser(resolve_entities=False, no_network=True)
            for name in names:
                if name.endswith(".rels"):
                    root = etree.fromstring(archive.read(name), parser)
                    for rel in root:
                        if rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink"):
                            raise ReportTemplateError("模板含外部资源链接，请将图片等资源嵌入文档后重试。")
    except (BadZipFile, KeyError) as exc:
        raise ReportTemplateError("无法读取 DOCX 文件，请用 Word 重新另存为 .docx 后导入。") from exc


def _level(paragraph):
    direct = paragraph._p.pPr
    outline = direct.find(qn("w:outlineLvl")) if direct is not None else None
    if outline is not None and int(outline.get(qn("w:val"), "9")) < 3:
        return int(outline.get(qn("w:val"))) + 1
    style = paragraph.style
    while style is not None:
        match = re.search(r"heading\s*([1-3])|标题\s*([1-3])", style.name, re.I)
        if match:
            return int(match[1] or match[2])
        ppr = style.element.find(qn("w:pPr"))
        outline = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        if outline is not None and int(outline.get(qn("w:val"), "9")) < 3:
            return int(outline.get(qn("w:val"))) + 1
        style = style.base_style
    return None


def _properties(paragraph):
    """Flatten style inheritance and the dominant text run's direct formatting."""
    ppr, rpr = OxmlElement("w:pPr"), OxmlElement("w:rPr")
    chain = []
    style = paragraph.style
    while style is not None and len(chain) < 30:
        chain.insert(0, style)
        style = style.base_style
    run = max((r for r in paragraph.runs if r.text.strip()), key=lambda r: len(r.text), default=None)
    for target, name in ((ppr, "pPr"), (rpr, "rPr")):
        defaults = paragraph.part.styles.element.find(qn("w:docDefaults"))
        default = defaults.find(f'{qn("w:" + name + "Default")}/{qn("w:" + name)}') if defaults is not None else None
        sources = [default] + [s.element.find(qn("w:" + name)) for s in chain]
        if name == "pPr":
            sources.append(paragraph._p.pPr)
        elif run is not None:
            sources.append(run._r.rPr)
        for source in sources:
            if source is None:
                continue
            for child in source:
                if child.tag in {qn("w:pStyle"), qn("w:sectPr"), qn("w:numPr")}:
                    continue
                previous = target.find(child.tag)
                if previous is not None and child.tag in {qn("w:rFonts"), qn("w:spacing"), qn("w:ind")}:
                    previous.attrib.update(child.attrib)
                else:
                    if previous is not None:
                        target.remove(previous)
                    target.append(deepcopy(child))
    return ppr, rpr


def _paragraph_roles(doc):
    samples, chapters = {}, []
    body_candidates = []
    references = False
    for index, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if not text or text == BODY_MARKER:
            continue
        level = _level(p)
        if text in {"参考文献", "References"}:
            references = True
        elif level == 1:
            references = False
        if "Title" == p.style.name or "{{REPORT_TITLE}}" in text:
            samples.setdefault("title", index)
        elif level:
            samples.setdefault(f"heading{level}", index)
            if level == 1 and text not in {"摘要", "目录", "参考文献", "Abstract", "References"}:
                chapters.append(re.sub(r"^(?:第[一二三四五六七八九十\d]+章\s*|\d+[.、\s]+)", "", text))
        elif references:
            samples.setdefault("reference", index)
        elif re.match(r"[图表]\s*\d|图题|表题", text):
            samples.setdefault("caption", index)
        elif re.match(r"关键词[：:]", text):
            samples.setdefault("keywords", index)
        elif len(text) > 15 and not re.search(r"\{\{|目录|课题名称|成文日期|生成机构", text):
            body_candidates.append(index)
    if body_candidates:
        # Pick the most-used paragraph style, then its longest representative.
        counts = Counter(doc.paragraphs[i].style.style_id for i in body_candidates)
        winning = counts.most_common(1)[0][0]
        samples["body"] = max((i for i in body_candidates if doc.paragraphs[i].style.style_id == winning),
                              key=lambda i: len(doc.paragraphs[i].text))
    return samples, normalize_chapters(chapters)


def _font_name(doc, fonts):
    if fonts is None:
        return None
    if fonts.get(qn("w:eastAsia")):
        return fonts.get(qn("w:eastAsia"))
    theme_name = fonts.get(qn("w:eastAsiaTheme"))
    if theme_name:
        from lxml import etree
        for relation in doc.part.rels.values():
            if relation.reltype.endswith('/theme'):
                theme = etree.fromstring(relation.target_part.blob)
                ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
                family = 'majorFont' if theme_name.startswith('major') else 'minorFont'
                node = theme.find(f'.//a:{family}', ns)
                if node is not None:
                    east = node.find('a:ea', ns)
                    if east is not None and east.get('typeface'):
                        return east.get('typeface')
                    chinese = node.find('a:font[@script="Hans"]', ns)
                    if chinese is not None:
                        return chinese.get('typeface')
    return fonts.get(qn('w:ascii'))


def _measure(doc, samples):
    profile = normalize_report_format_profile({})
    missing = []
    for role in ("body", "heading1", "heading2", "heading3", "caption", "reference"):
        if role not in samples:
            missing.append({"body": "正文示例", "heading1": "一级标题示例", "heading2": "二级标题示例",
                            "heading3": "三级标题示例", "caption": "图表题注示例", "reference": "参考文献示例"}[role])
            continue
        ppr, rpr = _properties(doc.paragraphs[samples[role]])
        fonts = rpr.find(qn("w:rFonts"))
        font = _font_name(doc, fonts)
        size = rpr.find(qn("w:sz"))
        if font:
            profile["font"]["heading" if role == "heading1" else role] = font
        if size is not None:
            profile["size"][role] = int(size.get(qn("w:val"))) / 2
        if role == "body":
            spacing = ppr.find(qn("w:spacing"))
            if spacing is not None and spacing.get(qn("w:line")):
                rule = spacing.get(qn("w:lineRule"), "auto")
                profile["layout"]["line_spacing_rule"] = rule
                profile["layout"]["line_spacing"] = int(spacing.get(qn("w:line"))) / (240 if rule == "auto" else 20)
    section = doc.sections[0]
    layout = profile["layout"]
    layout["page_width_mm"], layout["page_height_mm"] = round(section.page_width.mm, 2), round(section.page_height.mm, 2)
    layout["page"] = "Letter" if abs(section.page_width.mm - 215.9) < 2 else "A4"
    for side in ("top", "bottom", "left", "right"):
        layout[f"margin_{side}_mm"] = round(getattr(section, f"{side}_margin").mm, 2)
    layout["margin_mm"] = layout["margin_top_mm"]
    layout["include_toc"] = any(p.text.strip() == "目录" for p in doc.paragraphs) or bool(doc.element.xpath('.//w:instrText[contains(text(), "TOC")]'))
    start, cover = _body_start(doc)
    layout["cover"] = cover
    profile["chapters"] = []
    return profile, missing, start


def _body_start(doc):
    children = list(doc.element.body)
    marker = next((i for i, el in enumerate(children) if BODY_MARKER in "".join(el.itertext())), None)
    if marker is not None:
        return marker, marker > 0
    for p in doc.paragraphs:
        if p.text.strip() in {"摘要", "Abstract", "目录"} or _level(p) == 1:
            index = children.index(p._p)
            before = [x for x in doc.paragraphs if children.index(x._p) < index]
            has_cover = any(x.style.name == "Title" or "{{REPORT_TITLE}}" in x.text for x in before)
            return (index, True) if has_cover else (0, False)
    return 0, False


def import_template(name: str, data: bytes) -> dict:
    if Path(name).suffix.lower() != ".docx":
        raise ReportTemplateError("目前仅支持 .docx 模板，请先将文件另存为 Word 文档。")
    _validate_package(data)
    try:
        doc = Document(BytesIO(data))
        if len(doc.sections) != 1:
            raise ReportTemplateError("第一阶段支持单节模板；请合并分节并保留所需页眉页脚后导入。")
        samples, chapters = _paragraph_roles(doc)
        profile, missing, _ = _measure(doc, samples)
    except ReportTemplateError:
        raise
    except Exception as exc:
        raise ReportTemplateError("模板结构无法识别，请用 Word 另存为新的 DOCX 后重试。") from exc
    notices = []
    if missing:
        notices.append("以下项目未找到示例，将继承相关样式或系统默认值：" + "、".join(missing))
    if doc.element.xpath('.//w:txbxContent'):
        notices.append("模板包含文本框，第一阶段不保证复杂浮动排版完全一致。")
    if doc.sections[0]._sectPr.xpath('./w:cols[@w:num and @w:num!="1"]'):
        raise ReportTemplateError("第一阶段支持单栏模板，请改为单栏后重新导入。")
    template_id = sha256(data).hexdigest()[:24]
    info = {"id": template_id, "name": Path(name.replace("\\", "/")).name[:120],
            "created_at": datetime.now().isoformat(timespec="seconds"),
            "profile": profile, "chapters": chapters, "samples": samples, "notices": notices}
    TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
    path = TEMPLATE_DIR / f"{template_id}.docx"
    path.write_bytes(data)
    metadata = path.with_suffix(".json")
    temporary = metadata.with_suffix(".tmp")
    temporary.write_text(json.dumps(info, ensure_ascii=False, indent=2), encoding="utf-8")
    temporary.replace(metadata)
    return info


def resolve_template_profile(value) -> dict:
    """Trust the saved template values; apply only explicitly edited controls."""
    if isinstance(value, dict) and value.get("template"):
        candidate = value["template"]
        if not isinstance(candidate, dict):
            raise ReportTemplateError("报告模板配置无效，请重新选择模板。")
        template_path(candidate.get("id", ""))
    normalized = normalize_report_format_profile(value)
    template = normalized.get("template")
    if not template:
        return normalized
    template["overrides"] = [key for key in template["overrides"] if key in EDITABLE_PATHS]
    info = get_template(template["id"])
    merged = deepcopy(info["profile"])
    for key in template["overrides"]:
        if key not in EDITABLE_PATHS:
            continue
        group, field = key.split(".")
        merged[group][field] = normalized[group][field]
        if key == "layout.margin_mm":
            for side in ("top", "bottom", "left", "right"):
                merged["layout"][f"margin_{side}_mm"] = normalized[group][field]
        elif key == "layout.line_spacing":
            merged["layout"]["line_spacing_rule"] = "auto"
    merged["chapters"] = normalized["chapters"]
    merged["template"] = template
    return normalize_report_format_profile(merged)


def prepare_template_document(profile: dict, title: str):
    from .document_export import _font
    template = profile["template"]
    info = get_template(template["id"])
    doc = Document(template_path(template["id"]))
    samples = info["samples"]
    old_title = doc.paragraphs[samples["title"]].text if "title" in samples else ""
    properties = {role: _properties(doc.paragraphs[index]) for role, index in samples.items()}
    # Build canonical exporter styles before discarding sample text.
    for role, name in STYLE_ROLES.items():
        sample_index = samples.get(role, samples.get("body"))
        if role.startswith("heading"):
            sample_index = samples.get(role)
        if role == "title":
            sample_index = samples.get("title")
        source = doc.paragraphs[sample_index] if sample_index is not None else None
        target = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if source is not None:
            ppr, rpr = deepcopy(properties.get(role, properties.get("body")))
            for child in list(target.element):
                if child.tag in {qn("w:pPr"), qn("w:rPr")}:
                    target.element.remove(child)
            target.element.append(ppr)
            target.element.append(rpr)
        else:
            target.base_style = doc.styles["Normal"] if name != "Normal" else None
            _font(target.font, profile["font"].get("heading" if role.startswith("heading") else "body", "SimSun"),
                  profile["size"].get(role, profile["size"]["body"]), role.startswith("heading"))
        if role.startswith("heading"):
            outline = target.element.get_or_add_pPr().find(qn("w:outlineLvl"))
            if outline is None:
                outline = OxmlElement("w:outlineLvl")
                target.element.get_or_add_pPr().append(outline)
            outline.set(qn("w:val"), str(int(role[-1]) - 1))
            target.paragraph_format.keep_with_next = True
    for key in template["overrides"]:
        group, field = key.split(".")
        roles = ("body", "abstract", "keywords", "list") if field == "body" else ("heading1", "heading2", "heading3")
        if group == "font":
            for role in roles:
                style = doc.styles[STYLE_ROLES[role]]
                fonts = style.element.get_or_add_rPr().find(qn("w:rFonts"))
                if fonts is None:
                    fonts = OxmlElement("w:rFonts")
                    style.element.get_or_add_rPr().insert(0, fonts)
                fonts.set(qn("w:eastAsia"), profile[group][field])
                fonts.attrib.pop(qn("w:eastAsiaTheme"), None)
                if profile[group][field] in LATIN_REPORT_FONTS:
                    for attribute in ("ascii", "hAnsi", "cs"):
                        fonts.set(qn("w:" + attribute), profile[group][field])
                    for attribute in ("asciiTheme", "hAnsiTheme", "cstheme"):
                        fonts.attrib.pop(qn("w:" + attribute), None)
        elif group == "size":
            for role in (("body", "abstract", "keywords", "list") if field == "body" else (field,)):
                doc.styles[STYLE_ROLES[role]].font.size = Pt(profile[group][field])
        elif key == "layout.line_spacing":
            for role in ("body", "abstract", "keywords", "reference", "list"):
                doc.styles[STYLE_ROLES[role]].paragraph_format.line_spacing = profile[group][field]
    section = doc.sections[0]
    if "layout.margin_mm" in template["overrides"]:
        for side in ("top", "bottom", "left", "right"):
            setattr(section, side + "_margin", Mm(profile["layout"]["margin_mm"]))
    start, has_cover = _body_start(doc)
    keep_cover = profile["layout"]["cover"] and has_cover
    for index, element in enumerate(list(doc.element.body)):
        if element.tag != qn("w:sectPr") and (not keep_cover or index >= start):
            doc.element.body.remove(element)
    _replace_placeholders(doc, title, old_title)
    if keep_cover:
        # Ensure the generated abstract starts after the retained cover page.
        last = doc.paragraphs[-1] if doc.paragraphs else None
        if last is None or not last._p.xpath('.//w:br[@w:type="page"]'):
            doc.add_page_break()
    doc.core_properties.title = title
    doc.core_properties.author = ""
    update = doc.settings.element.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        doc.settings.element.append(update)
    update.set(qn("w:val"), "true")
    return doc, keep_cover


def _replace_placeholders(doc, title, old_title=""):
    # Work on text nodes to retain cover formatting and embedded drawings.
    replacements = {"{{REPORT_TITLE}}": title, "{{REPORT_DATE}}": datetime.now().strftime("%Y年%m月%d日")}
    if old_title.strip() and old_title != "{{REPORT_TITLE}}":
        replacements[old_title] = title
    roots = [doc.element]
    for rel in doc.part.rels.values():
        if not rel.is_external and rel.reltype.endswith(("/header", "/footer")):
            roots.append(rel.target_part.element)
    for root in roots:
        for p in root.iter(qn("w:p")):
            texts = list(p.iter(qn("w:t")))
            combined = "".join(t.text or "" for t in texts)
            replacement = combined
            for old, new in replacements.items():
                replacement = replacement.replace(old, new)
            pstyle = p.find("./" + qn("w:pPr") + "/" + qn("w:pStyle"))
            if pstyle is not None and pstyle.get(qn("w:val")) == "Title" and "{{" not in combined:
                replacement = title
            if replacement != combined and texts:
                texts[0].text = replacement
                for text in texts[1:]:
                    text.text = ""


def standard_template_bytes() -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25)
    from .document_export import _configure_document
    _configure_document(doc, "{{REPORT_TITLE}}")
    title = doc.add_paragraph("{{REPORT_TITLE}}", "Title")
    title.paragraph_format.space_before = Pt(100)
    doc.add_paragraph("成文日期：{{REPORT_DATE}}")
    doc.add_page_break()
    doc.add_paragraph(BODY_MARKER)
    doc.add_paragraph("摘要", "Heading 1")
    doc.add_paragraph("此段为正文排版示例，用于识别正文字体、字号、行距和缩进，生成报告时会被替换。")
    doc.add_paragraph("关键词：航空发动机；情报分析", "Report Keywords")
    doc.add_paragraph("目录", "Heading 1")
    for heading in ("研究背景与问题", "技术分析", "综合研判", "结论与建议"):
        doc.add_paragraph(heading, "Heading 1")
        doc.add_paragraph("本段为格式示例，实际正文由当前研究任务生成。")
    doc.add_paragraph("技术路线", "Heading 2")
    doc.add_paragraph("关键参数", "Heading 3")
    doc.add_paragraph("图1 排版示例", "Report Caption")
    doc.add_paragraph("参考文献", "Heading 1")
    doc.add_paragraph("[1] 作者. 文献题名[J]. 刊名, 年份, 卷(期): 页码.", "Report Reference")
    buffer = BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
