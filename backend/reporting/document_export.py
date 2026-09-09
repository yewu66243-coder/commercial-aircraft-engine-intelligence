"""Offline A4 Word/PDF exports sharing one structured Markdown parser.

The exporter formats supplied content; it never creates research claims,
headings, citation numbers, bibliography metadata, or figure/table numbers.
"""
from __future__ import annotations

from html import escape
import io
import math
import os
from pathlib import Path
import re
import sys
import unicodedata
from datetime import datetime
from urllib.parse import unquote, urlparse

import mistune

_INVALID_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]")
_ANCHOR = re.compile(r'<a\s+(?:id|name)=[\"\']([^\"\']+)[\"\']\s*>\s*</a>', re.I)
_CAPTION = re.compile(r"^(图|表)\s*\d+(?:[.－-]\d+)?(?:\s|[：:])")
_SOURCE = re.compile(r"^(?:来源|资料来源|图片来源|图源|数据来源|注|说明)\s*[：:]")
_FIGURE_SOURCE = re.compile(r"^(?:图片来源|图源)\s*[：:]")
_DLL_HANDLES = []


def clean_text(text: str) -> str:
    """Remove XML 1.0 forbidden characters, including malformed OCR surrogates."""
    return _INVALID_XML.sub("", str(text))


def _plain(tokens: list[dict]) -> str:
    return "".join(_plain(t["children"]) if "children" in t else
                   "\n" if t["type"] in {"softbreak", "linebreak"} else
                   "" if t["type"] in {"inline_html", "block_html"} else
                   t.get("raw", "") for t in tokens)


def _regular_weight(tokens):
    """Keep prose content and links while discarding inline bold emphasis."""
    result = []
    for token in tokens:
        if "children" in token:
            token = {**token, "children": _regular_weight(token["children"])}
        if token["type"] == "strong":
            result.extend(token.get("children", []))
        else:
            result.append(token)
    return result


def _blocks(text: str) -> list[dict]:
    tokens = mistune.create_markdown(renderer="ast", plugins=["table", "strikethrough"])(clean_text(text))
    result, anchors, section, title_seen = [], [], "body", False
    for token in tokens:
        kind = token["type"]
        if kind == "blank_line":
            continue
        raw = token.get("raw", "") if kind != "paragraph" else "".join(t.get("raw", "") for t in token.get("children", []))
        found = _ANCHOR.findall(raw)
        if found:
            anchors.extend(found)
            if not _plain(token.get("children", [])).strip():
                continue
            # Bibliography anchors commonly share their Markdown paragraph
            # with the reference text. Strip their markup and separator line.
            children = [child for child in token.get("children", []) if child["type"] != "inline_html"]
            while children and children[0]["type"] in {"softbreak", "linebreak"}:
                children.pop(0)
            token["children"] = children
        if kind == "block_html":
            continue
        token["anchors"], anchors = anchors, []
        plain = token["plain"] = _plain(token.get("children", [])).strip()
        if kind == "heading":
            level = token.get("attrs", {}).get("level", 2)
            if level == 1 and not title_seen:
                token["role"], title_seen = "title", True
            else:
                token["role"] = "heading"
                if level <= 2:
                    section = {"摘要": "abstract", "目录": "toc", "参考文献": "references"}.get(plain, "body")
                number = re.match(r"^(\d+(?:\.\d+)*)\s", plain)
                if number:
                    generated = "sec-" + number[1].replace(".", "-")
                    if generated not in token["anchors"]:
                        token["anchors"].append(generated)
        else:
            token["role"] = section
            if plain.startswith(("关键词", "关键字")):
                token["role"] = "keywords"
            elif _CAPTION.match(plain):
                token["role"] = "caption"
            elif _SOURCE.match(plain):
                token["role"] = "source"
                token["figure_source"] = bool(_FIGURE_SOURCE.match(plain))
            if kind == "paragraph" and any(t["type"] == "image" for t in token.get("children", [])):
                token["role"] = "figure"
            if token["role"] in {"body", "abstract"} and "children" in token:
                token["children"] = _regular_weight(token["children"])
        result.append(token)
    return result


def resolve_image_path(url: str, base_path: Path | str | None = None) -> Path | None:
    """Resolve supported local images without issuing network requests."""
    base = Path(base_path or Path.cwd()).resolve()
    parsed = urlparse(url)
    if parsed.scheme == "file":
        value = unquote(parsed.path)
        if os.name == "nt" and re.match(r"^/[A-Za-z]:", value):
            value = value[1:]
        path = Path(value)
    elif url.startswith("/outputs/"):
        path = base / unquote(parsed.path.lstrip("/"))
    elif not parsed.scheme and not parsed.netloc:
        path = base / unquote(parsed.path)
    else:
        return None
    return path.resolve() if path.is_file() else None


def _safe_link(url):
    if url.startswith("/api/local-library/"):
        # Downloaded documents have no browser origin. Point local source
        # links to this installation's API; hosted installs may set an origin.
        origin = os.environ.get("REPORT_EXPORT_BASE_URL", "http://127.0.0.1:8000").rstrip("/")
        parsed_origin = urlparse(origin)
        if parsed_origin.scheme in {"http", "https"} and parsed_origin.netloc:
            return origin + url
        return "http://127.0.0.1:8000" + url
    return url if urlparse(url).scheme.lower() in {"", "http", "https", "file", "mailto"} else ""


def _html_inline(tokens):
    parts = []
    for token in tokens:
        kind, children = token["type"], _html_inline(token.get("children", []))
        if kind in {"text", "codespan"}:
            parts.append(escape(token.get("raw", "")))
        elif kind in {"strong", "emphasis", "strikethrough"}:
            tag = {"strong": "strong", "emphasis": "em", "strikethrough": "s"}[kind]
            parts.append(f"<{tag}>{children}</{tag}>")
        elif kind == "link":
            href = _safe_link(token.get("attrs", {}).get("url", ""))
            parts.append(f'<a href="{escape(href, quote=True)}">{children}</a>' if href else children)
        elif kind in {"softbreak", "linebreak"}:
            parts.append("\n" if kind == "softbreak" else "<br>")
        elif kind not in {"inline_html", "image"}:
            parts.append(children)
    return "".join(parts)


def _table_rows(token):
    rows = []
    for child in token.get("children", []):
        if child["type"] == "table_head":
            rows.append(child["children"])
        elif child["type"] == "table_body":
            rows.extend(row["children"] for row in child["children"])
    return rows


def _column_widths(rows):
    count = max((len(row) for row in rows), default=1)
    weights = [max(5, min(30, max((len(_plain(row[col].get("children", [])))
               for row in rows if col < len(row)), default=5))) ** 0.6 for col in range(count)]
    return [weight / sum(weights) for weight in weights]


def _table_needs_page_breaks(rows, widths):
    """Conservatively estimate height in points at the PDF table's 10.5 pt.

    An oversized table must be allowed to flow immediately; applying `avoid`
    to it can otherwise strand its heading on a nearly empty preceding page.
    """
    height = 0
    for row in rows:
        row_lines = 1
        for cell, width in zip(row, widths):
            chars_per_line = max(1, ((160 * 72 / 25.4) * width - 8) / 5.25)
            text_lines = _plain(cell.get("children", [])).splitlines() or [""]
            wrapped = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(char) in {"W", "F"} else 1
                                             for char in line) / chars_per_line)) for line in text_lines)
            row_lines = max(row_lines, wrapped)
        height += row_lines * 10.5 * 1.25 + 8
    # Keep space for an adjacent heading/caption in the 700 pt A4 text area.
    return height > 620


def _list_html(items):
    return "".join("<li>" + "".join(_html_inline(child.get("children", []))
        if child["type"] != "list" else "<ul>" + _list_html(child["children"]) + "</ul>"
        for child in item.get("children", [])) + "</li>" for item in items)


def build_report_html(text: str, base_path: Path | str | None = None) -> str:
    """Build printable semantic HTML, including the supplied static contents."""
    blocks = _blocks(text)
    title = next((b["plain"] for b in blocks if b["role"] == "title"), "研究报告")
    parts = []
    for index, token in enumerate(blocks):
        kind, role = token["type"], token["role"]
        content = _html_inline(token.get("children", []))
        anchors = token.get("anchors", [])
        anchor_attr = f' id="{escape(anchors[0], quote=True)}"' if anchors else ""
        parts.extend(f'<a id="{escape(a, quote=True)}"></a>' for a in anchors[1:])
        if kind == "heading":
            level = min(6, token.get("attrs", {}).get("level", 2))
            attrs = f' class="{role}"' + anchor_attr
            if role == "title":
                attrs += f' data-short-title="{escape(title[:32], quote=True)}"'
            parts.append(f"<h{level}{attrs}>{content}</h{level}>")
        elif role == "figure":
            for item in token.get("children", []):
                if item["type"] != "image":
                    continue
                alt = _plain(item.get("children", [])).strip()
                path = resolve_image_path(item.get("attrs", {}).get("url", ""), base_path)
                visual = (f'<img src="{escape(path.as_uri(), quote=True)}" alt="{escape(alt, quote=True)}">'
                          if path else f'<p class="source">图片不可用：{escape(alt)}</p>')
                following = blocks[index + 1] if index + 1 < len(blocks) else {}
                duplicate = following.get("role") == "caption" and following.get("plain") == alt
                caption = f"<figcaption>{escape(alt)}</figcaption>" if alt and not duplicate else ""
                parts.append(f'<figure{anchor_attr}>{visual}{caption}</figure>')
                anchor_attr = ""
            surrounding = _html_inline([t for t in token.get("children", []) if t["type"] != "image"])
            if surrounding.strip():
                parts.append(f'<p class="body">{surrounding}</p>')
        elif kind == "table":
            rows = _table_rows(token)
            widths = _column_widths(rows)
            cols = "".join(f'<col style="width:{w * 100:.2f}%">' for w in widths)
            table_class = ' class="long-table"' if _table_needs_page_breaks(rows, widths) else ""
            parts.append(f'<table{anchor_attr}{table_class}><colgroup>{cols}</colgroup>')
            for n, row in enumerate(rows):
                group, tag = ("thead", "th") if n == 0 else ("tbody", "td")
                if n <= 1:
                    parts.append(f"<{group}>")
                parts.append("<tr>" + "".join(f'<{tag}>{_html_inline(c.get("children", []))}</{tag}>' for c in row) + "</tr>")
                if n == 0:
                    parts.append("</thead>")
            if len(rows) > 1:
                parts.append("</tbody>")
            parts.append("</table>")
        elif kind == "list":
            tag = "ol" if token.get("attrs", {}).get("ordered") else "ul"
            parts.append(f'<{tag} class="{role}"{anchor_attr}>{_list_html(token.get("children", []))}</{tag}>')
        elif kind == "block_code":
            parts.append(f'<pre{anchor_attr}>{escape(token.get("raw", ""))}</pre>')
        elif kind == "block_quote":
            parts.append(f"<blockquote{anchor_attr}>" + "".join(f'<p>{_html_inline(c.get("children", []))}</p>' for c in token.get("children", [])) + "</blockquote>")
        elif kind != "thematic_break":
            classes = role + (" figure-source" if token.get("figure_source") else "")
            parts.append(f'<p class="{classes}"{anchor_attr}>{content}</p>')
    return '<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8"><title>' + escape(title) + '</title></head><body>' + "\n".join(parts) + '</body></html>'


def _font(font, chinese="SimSun", size=12, bold=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    font.name, font.size, font.bold = "Times New Roman", Pt(size), bold
    font.color.rgb = RGBColor(0, 0, 0)
    rpr = font._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attribute in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + attribute), "Times New Roman")
    fonts.set(qn("w:eastAsia"), chinese)
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        fonts.attrib.pop(qn("w:" + attribute), None)


def _word_inline(paragraph, tokens, bold=False, italic=False):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    for token in tokens:
        kind = token["type"]
        if kind in {"strong", "emphasis", "strikethrough"}:
            _word_inline(paragraph, token.get("children", []), bold or kind == "strong", italic or kind == "emphasis")
        elif kind == "link":
            url = _safe_link(token.get("attrs", {}).get("url", ""))
            if not url:
                _word_inline(paragraph, token.get("children", []), bold, italic)
                continue
            hyperlink = OxmlElement("w:hyperlink")
            if url.startswith("#"):
                hyperlink.set(qn("w:anchor"), unquote(url[1:]))
            else:
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
            hyperlink.set(qn("w:history"), "1")
            run = paragraph.add_run(_plain(token.get("children", [])))
            run.bold, run.italic = bold, italic
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        elif kind in {"text", "codespan", "softbreak", "linebreak"}:
            # A Markdown source newline is flowing whitespace, not Shift+Enter.
            run = paragraph.add_run(" " if kind == "softbreak" else "\n" if kind == "linebreak" else token.get("raw", ""))
            run.bold, run.italic = bold, italic
        elif kind not in {"inline_html", "image"}:
            _word_inline(paragraph, token.get("children", []), bold, italic)


def _word_bookmarks(paragraph, anchors, known_anchors):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    for anchor in anchors:
        if anchor in known_anchors:
            continue
        known_anchors.add(anchor)
        bookmark_id = str(len(known_anchors))
        start, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        start.set(qn("w:id"), bookmark_id)
        start.set(qn("w:name"), anchor)
        end.set(qn("w:id"), bookmark_id)
        # w:pPr must remain the first child of a paragraph in valid OOXML.
        paragraph._p.insert(1 if paragraph._p.pPr is not None else 0, start)
        paragraph._p.append(end)


def _field(paragraph, instruction: str, placeholder: str = ""):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    if placeholder:
        text = OxmlElement("w:t")
        text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    if placeholder:
        run._r.append(text)
    run._r.append(end)
    return run


def _display_heading(doc, text: str, size: int = 15):
    from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
    from docx.shared import Pt
    paragraph = doc.add_paragraph()
    paragraph.alignment = Align.LEFT
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(12)
    run = paragraph.add_run(text)
    _font(run.font, "SimHei", size, True)
    return paragraph


def _add_cover(doc, title: str):
    from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
    from docx.shared import Pt
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    for _ in range(3):
        doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = Align.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    run = subtitle.add_run("商用航空发动机情报研究报告")
    _font(run.font, "SimHei", 16, True)

    heading = doc.add_paragraph(style="Title")
    heading.alignment = Align.CENTER
    heading.paragraph_format.first_line_indent = Pt(0)
    heading.paragraph_format.space_before = Pt(24)
    heading.paragraph_format.space_after = Pt(12)
    heading.add_run(title)

    sample = doc.add_paragraph()
    sample.alignment = Align.CENTER
    sample.paragraph_format.first_line_indent = Pt(0)
    sample.paragraph_format.space_before = Pt(18)
    run = sample.add_run("规范论文格式报告")
    _font(run.font, "SimSun", 14, False)

    for _ in range(5):
        doc.add_paragraph()
    for label, value in [
        ("课题名称", title),
        ("报告类型", "技术情报研究报告"),
        ("生成机构", "商用航空发动机情报工作台"),
        ("成文日期", datetime.now().strftime("%Y年%m月%d日")),
    ]:
        paragraph = doc.add_paragraph()
        paragraph.alignment = Align.CENTER
        paragraph.paragraph_format.first_line_indent = Pt(0)
        paragraph.paragraph_format.space_after = Pt(8)
        run = paragraph.add_run(f"{label}：{value}")
        _font(run.font, "SimSun", 12, False)
    doc.add_page_break()


def _add_word_toc(doc):
    from docx.shared import Pt
    _display_heading(doc, "目录", 15)
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    _field(paragraph, 'TOC \\o "1-2" \\h \\z \\u', "目录将在 Word 中自动更新")
    doc.add_page_break()


def _configure_document(doc, title):
    from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(25)
    section.header_distance = section.footer_distance = Mm(12.5)
    section.different_first_page_header_footer = True
    for grid in list(section._sectPr.findall(qn("w:docGrid"))):
        section._sectPr.remove(grid)
    # The bundled Word template contains a blue rule in Title. Eliminate
    # inherited paragraph rules while retaining the explicit table borders.
    for border in list(doc.styles._element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
    compat = doc.settings._element.find(qn("w:compat"))
    if compat is None:
        compat = OxmlElement("w:compat")
        doc.settings._element.append(compat)
    no_expand = compat.find(qn("w:doNotExpandShiftReturn"))
    if no_expand is None:
        no_expand = OxmlElement("w:doNotExpandShiftReturn")
        compat.insert(0, no_expand)
    no_expand.set(qn("w:val"), "1")
    update_fields = doc.settings._element.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        doc.settings._element.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    normal = doc.styles["Normal"]
    _font(normal.font)
    pf = normal.paragraph_format
    pf.alignment, pf.line_spacing = Align.JUSTIFY, 1.5
    pf.first_line_indent, pf.space_after = Pt(24), Pt(6)
    pf.widow_control = True
    snap = OxmlElement("w:snapToGrid")
    snap.set(qn("w:val"), "0")
    normal._element.get_or_add_pPr().insert_element_before(
        snap, "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange")
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12), ("Heading 4", 12), ("Heading 5", 12)):
        style = doc.styles[name]
        _font(style.font, "SimHei", size, True)
        sf = style.paragraph_format
        sf.first_line_indent, sf.line_spacing = Pt(0), 1.25
        sf.space_before, sf.space_after = Pt(0 if name == "Title" else 12), Pt(16 if name == "Title" else 6)
        sf.keep_with_next, sf.keep_together, sf.page_break_before = True, True, False
        sf.alignment = Align.CENTER if name == "Title" else Align.LEFT
    for name in ("Report Abstract", "Report Keywords", "Report Caption", "Report Source", "Report Reference", "Report Contents", "Report Table", "Report List"):
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        _font(style.font, size=10.5 if name in {"Report Caption", "Report Source", "Report Table"} else 12)
        sf = style.paragraph_format
        sf.first_line_indent, sf.space_after = Pt(0), Pt(4 if name == "Report Table" else 6)
        sf.line_spacing = 1.25 if name in {"Report Table", "Report Source"} else 1.5
        sf.alignment = Align.CENTER if name == "Report Caption" else Align.JUSTIFY if name in {"Report Abstract", "Report Keywords"} else Align.LEFT
    for name in ("Header", "Footer"):
        _font(doc.styles[name].font, size=9)
        doc.styles[name].paragraph_format.first_line_indent = Pt(0)
    header = section.header.paragraphs[0]
    header.text, header.alignment = title[:32], Align.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = Align.CENTER
    footer.add_run("第 ")
    _field(footer, " PAGE ", "1")
    footer.add_run(" 页")
    doc.core_properties.title, doc.core_properties.subject, doc.core_properties.author = title, "中文专题研究报告", ""


def _word_table(doc, token):
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt
    rows, widths = _table_rows(token), _column_widths(_table_rows(token))
    table = doc.add_table(rows=0, cols=len(widths))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Mm(160 * width)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + edge)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), "D9D9D9")
        borders.append(border)
    table._tbl.tblPr.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + side)
        node.set(qn("w:w"), "80")
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    table._tbl.tblPr.append(margins)
    for index, row_tokens in enumerate(rows):
        row = table.add_row()
        if index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for col, cell in enumerate(row.cells):
            cell.width = Mm(160 * widths[col])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index == 0 or index % 2 == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "D9EAF7" if index == 0 else "F7F9FB")
                cell._tc.get_or_add_tcPr().append(shade)
            p = cell.paragraphs[0]
            p.style = doc.styles["Report Table"]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.keep_with_next, p.paragraph_format.keep_together = index == 0, False
            if col < len(row_tokens):
                data = row_tokens[col]
                _word_inline(p, data.get("children", []), bold=index == 0)
                align = data.get("attrs", {}).get("align")
                if align == "center" or len(_plain(data.get("children", []))) <= 8:
                    p.alignment = Align.CENTER
                elif align == "right":
                    p.alignment = Align.RIGHT
            if index == 0:
                cb, bottom = OxmlElement("w:tcBorders"), OxmlElement("w:bottom")
                bottom.set(qn("w:val"), "single")
                bottom.set(qn("w:sz"), "6")
                bottom.set(qn("w:color"), "D9D9D9")
                cb.append(bottom)
                cell._tc.get_or_add_tcPr().append(cb)


def render_word(text: str, destination: str | Path, base_path: Path | str | None = None) -> None:
    """Write editable A4 DOCX; the async wrapper handles raised failures."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH as Align
    from docx.shared import Mm, Pt
    from PIL import Image
    blocks = _blocks(text)
    title = next((b["plain"] for b in blocks if b["role"] == "title"), "研究报告")
    doc = Document()
    _configure_document(doc, title)
    _add_cover(doc, title)
    known_anchors = set()
    styles = {"abstract": "Report Abstract", "keywords": "Report Keywords", "caption": "Report Caption", "source": "Report Source", "references": "Report Reference", "toc": "Report Contents"}
    skipping_toc = False
    for index, token in enumerate(blocks):
        kind, role = token["type"], token["role"]
        if kind == "heading":
            plain = token.get("plain", "")
            if role == "title":
                continue
            if plain == "目录":
                doc.add_page_break()
                _add_word_toc(doc)
                skipping_toc = True
                continue
            if skipping_toc:
                skipping_toc = False
            if plain == "参考文献":
                doc.add_page_break()
            style = "Title" if role == "title" else "Heading " + str(min(5, max(1, token.get("attrs", {}).get("level", 2) - 1)))
            paragraph = doc.add_paragraph(style=style)
            _word_inline(paragraph, token.get("children", []), bold=True)
            _word_bookmarks(paragraph, token.get("anchors", []), known_anchors)
        elif skipping_toc:
            continue
        elif kind == "table":
            _word_table(doc, token)
        elif role == "figure":
            for item in token.get("children", []):
                if item["type"] != "image":
                    continue
                alt = _plain(item.get("children", [])).strip()
                path = resolve_image_path(item.get("attrs", {}).get("url", ""), base_path)
                paragraph = doc.add_paragraph(style="Report Caption")
                paragraph.paragraph_format.keep_with_next = bool(alt)
                paragraph.paragraph_format.first_line_indent = Pt(0)
                if path:
                    with Image.open(path) as original:
                        source, buffer = original.convert("RGB"), io.BytesIO()
                        source.save(buffer, format="PNG")
                        buffer.seek(0)
                        width, height = source.size
                        scale = min(160 / width, 175 / height)
                        shape = paragraph.add_run().add_picture(buffer, width=Mm(width * scale), height=Mm(height * scale))
                        shape._inline.docPr.set("descr", alt)
                else:
                    paragraph.add_run("图片不可用：" + alt)
                following = blocks[index + 1] if index + 1 < len(blocks) else {}
                duplicate = following.get("role") == "caption" and following.get("plain") == alt
                if alt and not duplicate:
                    caption = doc.add_paragraph(alt, style="Report Caption")
                    caption.paragraph_format.first_line_indent = Pt(0)
                    caption.paragraph_format.keep_with_next = following.get("role") == "source"
            surrounding = [t for t in token.get("children", []) if t["type"] != "image"]
            if _plain(surrounding).strip():
                _word_inline(doc.add_paragraph(), surrounding)
        elif kind == "list":
            def add_list(items, depth=0, ordered=False, start=1):
                for number, item in enumerate(items, start):
                    for child in item.get("children", []):
                        if child["type"] == "list":
                            add_list(child["children"], depth + 1, child.get("attrs", {}).get("ordered", False))
                            continue
                        p = doc.add_paragraph(style="Report Contents" if role == "toc" else "Report List")
                        p.paragraph_format.left_indent = Pt(12 * depth if role == "toc" else 18 + 18 * depth)
                        p.paragraph_format.first_line_indent = Pt(0 if role == "toc" else -12)
                        if role != "toc":
                            p.add_run(str(number) + ". " if ordered else "• ")
                        _word_inline(p, child.get("children", []))
            add_list(token.get("children", []), ordered=token.get("attrs", {}).get("ordered", False), start=token.get("attrs", {}).get("start", 1))
        elif kind == "block_quote":
            for child in token.get("children", []):
                paragraph = doc.add_paragraph(style="Report Abstract")
                paragraph.paragraph_format.left_indent = Pt(24)
                _word_inline(paragraph, child.get("children", []))
        elif kind == "block_code":
            doc.add_paragraph(token.get("raw", "").rstrip(), style="Report Source")
        elif kind != "thematic_break":
            paragraph = doc.add_paragraph(style=styles.get(role, "Normal"))
            _word_inline(paragraph, token.get("children", []))
            _word_bookmarks(paragraph, token.get("anchors", []), known_anchors)
            if role == "keywords" and index + 1 < len(blocks) and blocks[index + 1].get("plain") == "目录":
                paragraph.paragraph_format.keep_with_next = False
            if token.get("figure_source"):
                paragraph.alignment = Align.CENTER
            elif role == "references":
                paragraph.paragraph_format.left_indent, paragraph.paragraph_format.first_line_indent = Pt(24), Pt(-24)
                paragraph.paragraph_format.keep_together = False
            elif role == "caption":
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.keep_with_next = token["plain"].startswith("表") or (index + 1 < len(blocks) and blocks[index + 1]["role"] == "source")
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def render_pdf(text: str, destination: str | Path, base_path: Path | str | None = None) -> None:
    """Use module-relative GTK and CSS paths, independent of the working dir."""
    project = Path(__file__).resolve().parents[2]
    if sys.platform == "win32" and hasattr(os, "add_dll_directory") and not _DLL_HANDLES:
        gtk_path = project / "dependencies" / "gtk3" / "bin"
        if gtk_path.is_dir():
            _DLL_HANDLES.append(os.add_dll_directory(str(gtk_path)))
    from weasyprint import CSS, HTML
    css_path = Path(__file__).resolve().parents[1] / "styles" / "pdf_styles.css"
    path = Path(destination)
    path.parent.mkdir(parents=True, exist_ok=True)
    HTML(string=build_report_html(text, base_path), base_url=str(Path(base_path or Path.cwd()).resolve())).write_pdf(str(path), stylesheets=[CSS(filename=str(css_path))])
