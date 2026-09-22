"""Validation and storage for entity-evaluation ground-truth uploads."""

from __future__ import annotations

import hashlib
import json
import math
import os
import re
import tempfile
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


MAX_UPLOAD_BYTES = 5 * 1024 * 1024
ALLOWED_SUFFIXES = {".json", ".xlsx"}

_CATEGORY_ALIASES = {
    "organization": {"机构", "企业", "公司", "制造商", "监管机构", "研究机构", "organization"},
    "model": {"型号", "产品", "发动机型号", "部件型号", "平台", "model"},
    "material": {"材料", "合金", "涂层", "复合材料", "工艺材料", "material"},
    "parameter": {"参数", "性能参数", "技术指标", "数值", "规格", "parameter"},
    "time": {"时间", "日期", "年份", "阶段", "里程碑", "time"},
}
_REQUIRED_XLSX_COLUMNS = ("类别", "名称")
_OPTIONAL_XLSX_COLUMNS = ("别名", "数值", "单位", "容差")


class GroundTruthValidationError(ValueError):
    """A safe, stable validation error suitable for API responses."""

    def __init__(self, code: str, message: str):
        self.code = code
        self.message = message
        super().__init__(message)


def _error(code: str, message: str) -> None:
    raise GroundTruthValidationError(code, message)


def _normalize_category(value: Any) -> str:
    normalized = str(value or "").strip().lower()
    for category, aliases in _CATEGORY_ALIASES.items():
        if normalized in {alias.lower() for alias in aliases}:
            return category
    return "other"


def _normalized_key(value: Any) -> str:
    return re.sub(r"[\s\-_./:：,，;；()（）\[\]【】'\"“”‘’]+", "", str(value or "").lower())


def _clean_aliases(value: Any, location: str = "") -> list[str]:
    if value is None:
        return []
    if not isinstance(value, list) or not all(isinstance(alias, str) for alias in value):
        _error("invalid_aliases", f"标准答案实体别名必须是字符串列表。{location}")
    aliases: list[str] = []
    seen: set[str] = set()
    for alias in value:
        cleaned = _normalized_key(alias)
        if cleaned and cleaned not in seen:
            seen.add(cleaned)
            aliases.append(cleaned)
    return aliases


def _validate_parameter_fields(
    entity: dict[str, Any], normalized: dict[str, Any], location: str = ""
) -> None:
    if "value" in entity and entity["value"] is not None:
        value = entity["value"]
        if isinstance(value, bool) or not isinstance(value, (int, float, str)):
            _error("invalid_parameter", f"标准答案参数值必须是数值或数值字符串。{location}")
        if isinstance(value, float) and not math.isfinite(value):
            _error("invalid_parameter", f"标准答案参数值必须是有限数值。{location}")
        normalized["value"] = value.strip() if isinstance(value, str) else value
    if "unit" in entity and entity["unit"] is not None:
        unit = entity["unit"]
        if not isinstance(unit, str):
            _error("invalid_parameter", f"标准答案参数的单位必须是字符串。{location}")
        normalized["unit"] = unit.strip()

    tolerance = entity.get("tolerance", 0.01)
    if (
        isinstance(tolerance, bool)
        or not isinstance(tolerance, (int, float))
        or not math.isfinite(float(tolerance))
        or float(tolerance) < 0
    ):
        _error("invalid_parameter", f"标准答案参数的容差必须是非负有限数值。{location}")
    normalized["tolerance"] = float(tolerance)


def canonicalize_ground_truth_payload(payload: Any, expected_task: str) -> dict[str, Any]:
    """Validate a decoded payload and return the sole canonical representation."""
    if not isinstance(payload, dict):
        _error("invalid_schema", "标准答案顶层结构必须是对象。")

    declared_task = payload.get("task")
    if declared_task is None or declared_task == "":
        task = expected_task
    elif not isinstance(declared_task, str) or declared_task != expected_task:
        _error("task_mismatch", "标准答案声明的任务与当前任务不一致。")
    else:
        task = declared_task

    entities = payload.get("entities")
    if not isinstance(entities, list):
        _error("invalid_schema", "标准答案实体列表结构无效。")
    if not entities:
        _error("empty_entities", "标准答案实体列表不能为空。")

    normalized_entities: list[dict[str, Any]] = []
    duplicate_keys: set[tuple[str, str, str, str]] = set()
    for entity in entities:
        if not isinstance(entity, dict):
            _error("invalid_entity", "标准答案包含无效实体。")
        row_context = entity.get("_row_context")
        location = f"（{row_context}）" if isinstance(row_context, str) and row_context else ""
        raw_type = entity.get("type")
        if raw_type is None or (isinstance(raw_type, str) and not raw_type.strip()):
            raw_type = entity.get("category")
        name = entity.get("name")
        if not isinstance(raw_type, str) or not raw_type.strip() or not isinstance(name, str) or not name.strip():
            _error("invalid_entity", f"标准答案实体缺少有效类别或名称。{location}")

        normalized = {
            "type": _normalize_category(raw_type),
            "name": name.strip(),
            "aliases": _clean_aliases(entity.get("aliases", []), location),
        }
        if normalized["type"] == "parameter":
            _validate_parameter_fields(entity, normalized, location)

        duplicate_key = (
            normalized["type"],
            _normalized_key(normalized["name"]),
            _normalized_key(normalized.get("value", "")),
            _normalized_key(normalized.get("unit", "")),
        )
        if duplicate_key in duplicate_keys:
            _error("duplicate_entity", f"标准答案包含重复实体。{location}")
        duplicate_keys.add(duplicate_key)
        normalized_entities.append(normalized)

    return {"task": task, "entities": normalized_entities}


def _validate_upload_path(path: Path) -> None:
    suffix = path.suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        _error("unsupported_file_type", "仅支持 JSON 或 XLSX 标准答案文件。")
    try:
        size = path.stat().st_size
    except OSError:
        _error("unreadable_file", "标准答案文件无法读取。")
    if size == 0:
        _error("empty_file", "标准答案文件不能为空。")
    if size > MAX_UPLOAD_BYTES:
        _error("file_too_large", "标准答案文件不能超过 5 MiB。")


def _formula_cells_without_cached_values(path: Path, worksheet_index: int) -> set[str]:
    """Find formula coordinates which have no cached value in the worksheet XML."""
    try:
        with zipfile.ZipFile(path) as archive:
            root = ElementTree.fromstring(archive.read(f"xl/worksheets/sheet{worksheet_index}.xml"))
    except (KeyError, OSError, zipfile.BadZipFile, ElementTree.ParseError):
        return set()

    namespace = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    missing: set[str] = set()
    for cell in root.iter(f"{namespace}c"):
        formula = cell.find(f"{namespace}f")
        cached = cell.find(f"{namespace}v")
        if formula is not None and (cached is None or cached.text is None):
            coordinate = cell.get("r")
            if coordinate:
                missing.add(coordinate)
    return missing


def _split_aliases(value: Any) -> list[str]:
    if value is None:
        return []
    return [part.strip() for part in re.split(r"[;；]", str(value)) if part.strip()]


def _xlsx_payload(path: Path) -> dict[str, Any]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    except Exception:
        _error("invalid_excel", "标准答案 XLSX 文件无法读取。")

    try:
        visible_sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        if not visible_sheets:
            _error("invalid_excel", "标准答案 XLSX 文件没有可见工作表。")
        worksheet = visible_sheets[0]
        worksheet_index = workbook.worksheets.index(worksheet) + 1
        formula_cells = _formula_cells_without_cached_values(path, worksheet_index)
        rows = worksheet.iter_rows(values_only=False)
        try:
            header_cells = next(rows)
        except StopIteration:
            _error("missing_columns", f"工作表{worksheet.title}缺少必需列。")

        headers = {
            str(cell.value).strip(): index
            for index, cell in enumerate(header_cells)
            if cell.value is not None and str(cell.value).strip()
        }
        missing_columns = [column for column in _REQUIRED_XLSX_COLUMNS if column not in headers]
        if missing_columns:
            _error("missing_columns", f"工作表{worksheet.title}缺少必需列：{'、'.join(missing_columns)}。")

        entities: list[dict[str, Any]] = []
        known_columns = (*_REQUIRED_XLSX_COLUMNS, *_OPTIONAL_XLSX_COLUMNS)
        for row_number, row_cells in enumerate(rows, start=2):
            if all(cell.value is None for cell in row_cells):
                continue
            for column_index, _cell in enumerate(row_cells, start=1):
                if f"{get_column_letter(column_index)}{row_number}" in formula_cells:
                    _error(
                        "formula_without_cached_value",
                        f"工作表{worksheet.title}，第 {row_number} 行的公式没有缓存值。",
                    )
            values = {
                column: row_cells[index].value if index < len(row_cells) else None
                for column, index in headers.items()
                if column in known_columns
            }
            entity: dict[str, Any] = {
                "type": values.get("类别"),
                "name": values.get("名称"),
                "aliases": _split_aliases(values.get("别名")),
                "_row_context": f"工作表{worksheet.title}，第 {row_number} 行",
            }
            for source, target in (("数值", "value"), ("单位", "unit"), ("容差", "tolerance")):
                if values.get(source) is not None:
                    entity[target] = values[source]
            entities.append(entity)
        return {"entities": entities}
    finally:
        workbook.close()


def load_ground_truth_upload(path: str | Path, expected_task: str) -> dict[str, Any]:
    """Load a JSON or XLSX upload and return its validated canonical form."""
    upload_path = Path(path)
    _validate_upload_path(upload_path)
    if upload_path.suffix.lower() == ".json":
        try:
            with upload_path.open("r", encoding="utf-8") as handle:
                payload = json.load(handle)
        except (OSError, UnicodeError, json.JSONDecodeError):
            _error("invalid_json", "标准答案文件不是有效 JSON。")
    else:
        payload = _xlsx_payload(upload_path)
    return canonicalize_ground_truth_payload(payload, expected_task)


def ground_truth_path_for_task(task: str, destination_dir: str | Path) -> Path:
    digest = hashlib.sha256(task.encode("utf-8")).hexdigest()[:20]
    return Path(destination_dir) / f"{digest}.json"


def persist_ground_truth_upload(
    content: bytes, original_name: str, task: str, destination_dir: str | Path
) -> dict[str, Any]:
    """Validate bytes and atomically store canonical JSON without trusting its filename."""
    if not isinstance(content, bytes):
        _error("invalid_upload", "标准答案上传内容无效。")
    suffix = Path(original_name or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        _error("unsupported_file_type", "仅支持 JSON 或 XLSX 标准答案文件。")
    if not content:
        _error("empty_file", "标准答案文件不能为空。")
    if len(content) > MAX_UPLOAD_BYTES:
        _error("file_too_large", "标准答案文件不能超过 5 MiB。")

    directory = Path(destination_dir)
    directory.mkdir(parents=True, exist_ok=True)
    temporary_paths: set[Path] = set()
    canonical_path = ground_truth_path_for_task(task, directory)
    try:
        with tempfile.NamedTemporaryFile(
            mode="wb", suffix=suffix, prefix=".ground-truth-", dir=directory, delete=False
        ) as temporary:
            temporary.write(content)
            upload_path = Path(temporary.name)
            temporary_paths.add(upload_path)
        canonical = load_ground_truth_upload(upload_path, task)

        with tempfile.NamedTemporaryFile(
            mode="w", encoding="utf-8", suffix=".tmp", prefix=".ground-truth-", dir=directory, delete=False
        ) as temporary:
            json.dump(canonical, temporary, ensure_ascii=False, separators=(",", ":"))
            canonical_temporary_path = Path(temporary.name)
            temporary_paths.add(canonical_temporary_path)
        os.replace(canonical_temporary_path, canonical_path)
        temporary_paths.remove(canonical_temporary_path)
    finally:
        for temporary_path in temporary_paths:
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass

    category_counts = dict(Counter(entity["type"] for entity in canonical["entities"]))
    return {
        "stored_name": canonical_path.name,
        "sha256": hashlib.sha256(content).hexdigest(),
        "entity_count": len(canonical["entities"]),
        "category_counts": category_counts,
    }
