"""Project-runtime XLSX interchange, sharing the ground-truth API's openpyxl.

The server manifest, not spreadsheet formulas or cell protection, is authoritative.
"""
from io import BytesIO
from zipfile import ZipFile
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_ROWS = 20000
LEGACY_HEADERS = ['编号', '类别', '实体或引用', '报告原句', '引用编号', '来源文件或URL',
           '原文位置', '候选证据摘录（非判定）', '材料入口', '人工判定', '判定理由', '核验原文摘录', '核验来源及位置']
HEADERS = LEGACY_HEADERS[:6] + ['正文位置 / 来源页码', '引用原文摘录（辅助核对）'] + LEGACY_HEADERS[8:]
FACT_FIELDS = ['attribute', 'value', 'unit', 'raw_value', 'conditions', 'statement_type',
               'association', 'citation_assignment', 'review_note', 'parent_claim']
FACT_HEADERS = ['核验属性', '数值', '完整单位', '原始数值表达', '时间/条件/限定语', '语句性质',
                '对象关联方式', '引用对应方式', '需人工确认事项', '完整上文原句']
V3_HEADERS = HEADERS + FACT_HEADERS
AUDIT_HEADERS = ['编号', '语句性质', '正文原句', '正文位置', '提取说明']
DECISIONS = ('正确', '错误', '待核验', '排除')
META_INPUTS = {'核验人', '复核人', '核验日期', '全量检查确认', '提取复核确认', '排除项复核确认', '异常或争议说明'}


class ManualReviewError(ValueError):
    pass


def text(value):
    return '' if value is None else str(value)


def _literal(value):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text(value))


def fixed_values(row):
    evidence = row.get('evidence') or []
    sources = row.get('source', '\n'.join(dict.fromkeys(text(e.get('source')) for e in evidence)))
    if row.get('body_location'):
        positions = [row['body_location']] + list(dict.fromkeys(
            f"引用{e.get('ref', '')} 第{e.get('page')}页" if e.get('page') is not None
            else '引用原文未记录可靠页码（见材料入口）' for e in evidence))
    else:
        # Preserve the fixed-cell signature of previously exported v1 workbooks.
        positions = [f"{e.get('source', '')} 第{e.get('page')}页" if e.get('page') is not None
                     else f"{e.get('source', '')}（无可靠页码）" for e in evidence]
    values = [row['id'], row.get('category', ''), row.get('name', ''), row['claim'],
            ' '.join(row.get('refs', [])), sources,
            '\n'.join(positions),
            '\n\n'.join(text(e.get('text')) for e in evidence)[:24000],
            f"evidence.html#{row['id']}"]
    return [_literal(value) for value in values]


def fact_values(row):
    return [_literal(row.get(field, '')) for field in FACT_FIELDS]


def audit_values(item, index):
    return [_literal(value) for value in [f'U{index}', item.get('statement_type', ''), item['claim'],
                                          item.get('body_location', ''), item['reason']]]


def _write(ws, values):
    ws.append(values)
    for cell in ws[ws.max_row]:
        if isinstance(cell.value, str):
            cell.value = _literal(cell.value)
            cell.data_type = 's'


def _style(ws, editable_start=0):
    # Copying sources must remain unrestricted; import validates fixed fields.
    ws.protection.sheet = False
    ws.freeze_panes = 'D2' if ws.max_column > 3 else 'A2'
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = '1:1'
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            editable = bool(editable_start and cell.column >= editable_start and cell.row > 1
                            and (editable_start == 1 or cell.column <= 13))
            cell.font = Font(name='微软雅黑', size=11, color='156082' if editable else '222222')
            cell.fill = PatternFill('solid', fgColor='EAF5FB' if editable else 'FFFFFF')
            cell.protection = Protection(locked=not editable)
    for cell in ws[1]:
        cell.font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='205867')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 24
    ws.row_dimensions[1].height = 32


def write_workbook(manifest, path):
    wb = Workbook()
    structured = manifest.get('extraction_schema') == 'object-attribute-v3'
    headers = V3_HEADERS if structured else HEADERS
    meta = wb.active
    meta.title = '任务与规则'
    _write(meta, ['项目', '内容'])
    for key, value in {
        '核验包ID': manifest['package_id'], '运行编号': manifest['run_id'],
        '报告SHA256': manifest['report_sha256'], '任务': manifest['task'],
        '报告题目': manifest.get('report_title', ''), '模型': manifest.get('model', ''),
        '提取规则版本': manifest.get('policy_version', ''),
        '导出清理': ('导出前删除“来源文件或URL”为空的整行，并移除名称/原句为空或明显乱码的无效行及空白/乱码摘录；'
                 '来源非空而摘录为空的行仍保留。不另存删除明细，不改动报告原文。'),
        '导出清理统计': ('移除无效行：{removed_rows}；移除无效摘录：{removed_excerpts}；'
                     '其中缺来源行：{missing_source_rows}；保留缺可读摘录项：{missing_evidence_rows}。'.format(
                         **{key: (manifest.get('export_cleanup') or {}).get(key, 0)
                            for key in ('removed_rows', 'removed_excerpts', 'missing_evidence_rows', 'missing_source_rows')})),
        '缺来源项口径': ('缺来源行从主表移除不代表正文对应事实已核验。'
                    '发生此类清理的核验包只统计保留样本，不据此认定全报告达标；'
                    '完整验收需修正报告引用后以新的报告版本重新核验。'),
        '提取范围': '最终保存的报告正文，含摘要及正文表格；不纳入检索候选、目录、关键词、参考文献列表或内部核验附件。',
        '正文公开网址数（去重）': manifest.get('public_url_count', 0),
        '开始时间': manifest.get('started_at', ''), '报告导出完成时间': manifest.get('completed_at', ''),
        '生成耗时（秒）': manifest.get('duration_seconds'),
        '指标要求': '耗时≤1200秒；实体准确率≥90%；公开引用准确率≥98%；至少10轮且每轮均达标。',
        '计分规则': '正确数/纳入核验总数；未核验不判达标；零样本不计100%；不采用半分。',
        '实体口径': '按对象—属性—数值核验，保留完整名称、单位与原句。同一句中对象与参数不能可靠关联时保留待确认，不猜测。不从候选材料添加实体；没有明确引用时来源留空。',
        '结构化字段': 'N列起为只读辅助字段；属性/对象待确认的行仍在分母中，须人工核对原句。预测、计划与实际发生不得合并。提取规则不能判断事实真伪。',
        '溯源口径': '仅正文实际使用的公开URL：直接链接或正文引用编号明确对应的网址；仅列在参考文献但未被正文引用的不纳入。同一URL用于不同原句分别核验。网页能打开不等于引用正确。',
        '填写方式': '蓝色为输入；J至M列填写；核验来源填写原文及页码/段落。片段不足请打开材料入口看完整已存原文。',
        '复制查证': '工作表不启用保护，可选择、复制来源与摘录。原始字段仍由后台校验；请勿修改编号、原句、来源及结构化记录。',
        '补充规则': '漏项与争议中编号填A开头唯一编号；类别填entity:参数或source:公开引用等；必须附正文原句。',
        '排除规则': '仅限重复或不属于核验范围，必须说明原因并由不同复核人确认。不得用排除移除错误或找不到依据的事实。',
        '全量检查': '逐段核对报告，补齐自动规则遗漏项；完成后在全量检查确认填写“确认”，不是抽样核验。',
        '提取审计': '检查“提取审计”表中的未识别、建议、研究问题和范围说明；如误分或漏抽事实，在“漏项与争议”补齐，再在提取复核确认填写“确认”。此表不是免责依据。',
        '标准答案': '本表不测漏抽召回率/F1。若验收要求，应另建固定资料集的人工标准答案。',
        '结果使用': 'Excel汇总为参考；导入时服务器重新计算并校验版本、行数和内容。报告修改后需重新导出核验包。',
        '核验人': '', '复核人': '', '核验日期': '', '全量检查确认': '', '提取复核确认': '',
        '排除项复核确认': '', '异常或争议说明': '',
    }.items():
        _write(meta, [key, value])
    _style(meta)
    meta.column_dimensions['A'].width = 25
    meta.column_dimensions['B'].width = 110
    for row in meta.iter_rows(min_row=2):
        meta.row_dimensions[row[0].row].height = 42
        if row[0].value in META_INPUTS:
            row[1].fill = PatternFill('solid', fgColor='EAF5FB')
            row[1].protection = Protection(locked=False)
    for title, kind in [('实体核验', 'entity'), ('溯源核验', 'source'), ('漏项与争议', None)]:
        ws = wb.create_sheet(title)
        _write(ws, headers)
        for row in manifest['rows']:
            if row['kind'] == kind:
                _write(ws, fixed_values(row) + ['待核验', '', '', ''] + (fact_values(row) if structured else []))
        if kind is None:
            for _ in range(30):
                _write(ws, [''] * len(headers))
        _style(ws, 10 if kind else 1)
        for letter, width in {'A': 12, 'D': 58, 'F': 42, 'H': 70, 'I': 24, 'K': 45, 'L': 60, 'M': 45}.items():
            ws.column_dimensions[letter].width = width
        ws.column_dimensions['G'].width = 32
        if kind == 'source':
            ws.column_dimensions['C'].width = 42
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height = 100
            if kind:
                row[8].hyperlink = row[8].value
        if kind:
            ws.protection.autoFilter = False
        dv = DataValidation(type='list', formula1='"' + ','.join(DECISIONS) + '"', allow_blank=True)
        dv.errorTitle, dv.error = '判定值无效', '请选择正确、错误、待核验或排除'
        dv.showErrorMessage = True
        ws.add_data_validation(dv)
        dv.add(f'J2:J{MAX_ROWS}')
    if structured:
        audit = wb.create_sheet('提取审计')
        _write(audit, AUDIT_HEADERS)
        for index, item in enumerate(manifest.get('extraction_audit', []), 1):
            _write(audit, audit_values(item, index))
        _style(audit)
        for column, width in {'A': 12, 'B': 25, 'C': 95, 'D': 45, 'E': 65}.items():
            audit.column_dimensions[column].width = width
        for index in range(2, audit.max_row + 1):
            audit.row_dimensions[index].height = 110
    summary = wb.create_sheet('指标汇总')
    _write(summary, ['指标', '纳入总数', '正确', '错误', '待核验', '排除', '已判定项正确率（非最终验收）'])
    for index, (sheet, kind) in enumerate([('实体核验', 'entity'), ('溯源核验', 'source')], 2):
        _write(summary, [sheet, '', '', '', '', '', ''])
        for column, decision in [('C', '正确'), ('D', '错误'), ('F', '排除')]:
            summary[f'{column}{index}'] = (f'=COUNTIF(\'{sheet}\'!J2:J{MAX_ROWS},"{decision}")+'
                f'COUNTIFS(\'漏项与争议\'!B2:B{MAX_ROWS},"{kind}:*",\'漏项与争议\'!J2:J{MAX_ROWS},"{decision}")')
        summary[f'B{index}'] = (f'=COUNTA(\'{sheet}\'!A2:A{MAX_ROWS})+'
            f'COUNTIF(\'漏项与争议\'!B2:B{MAX_ROWS},"{kind}:*")-F{index}')
        summary[f'E{index}'] = f'=B{index}-C{index}-D{index}'
        summary[f'G{index}'] = f'=IF(C{index}+D{index}=0,"待核验",C{index}/(C{index}+D{index}))'
        summary[f'G{index}'].number_format = '0.00%'
    _style(summary)
    summary.column_dimensions['G'].width = 42
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    wb.close()


def read_workbook(content):
    if len(content) > MAX_UPLOAD_BYTES:
        raise ManualReviewError('核验表不能超过10 MiB。')
    try:
        with ZipFile(BytesIO(content)) as z:
            entries = z.infolist()
            if len(entries) > 2000 or sum(e.file_size for e in entries) > 80 * 1024 * 1024:
                raise ManualReviewError('Excel解压体积过大。')
            if any('vbaproject' in e.filename.lower() or 'externallinks/' in e.filename.lower() for e in entries):
                raise ManualReviewError('不接受宏或外部工作簿链接。')
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except ManualReviewError:
        raise
    except Exception as exc:
        raise ManualReviewError('无法读取核验Excel。') from exc
    try:
        required = {'任务与规则', '实体核验', '溯源核验', '漏项与争议', '指标汇总'}
        if set(wb.sheetnames) not in (required, required | {'提取审计'}):
            raise ManualReviewError('工作表缺失或名称发生变化。')
        result = {}
        for name in set(wb.sheetnames) - {'指标汇总'}:
            ws = wb[name]
            if ws.max_row > MAX_ROWS or ws.max_column > len(V3_HEADERS):
                raise ManualReviewError('核验表行数或列数超过上限。')
            rows = []
            for row in ws:
                if any(c.data_type == 'f' for c in row):
                    raise ManualReviewError('输入区域不接受公式，请填写文本判定。')
                rows.append([text(c.value) for c in row])
            result[name] = rows
        return result
    finally:
        wb.close()
