import json
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document

from backend.reporting.citations import _format_gbt_reference
from backend.reporting.reference_metadata import read_local_metadata


ROOT = Path(__file__).resolve().parents[1]
RECORDS = ROOT / "outputs" / "records" / "evaluation_records.json"
SOURCE = ROOT / "outputs" / "航空发动机剩余寿命预测方法：技术路线、适航约束与运营影响研究_f792.docx"
OUTPUT = ROOT / "outputs" / "航空发动机剩余寿命预测方法：技术路线、适航约束与运营影响研究_f792_参考文献修订版.docx"


def enriched(item):
    record = dict(item)
    local = read_local_metadata(record)
    for key, value in local.items():
        if not value:
            continue
        if key == "metadata_provenance":
            record[key] = {**(record.get(key) or {}), **value}
        elif key == "source_type" and "学位" in str(value):
            record[key] = value
        elif not record.get(key):
            record[key] = value
    return record


def main():
    records = json.loads(RECORDS.read_text(encoding="utf-8"))
    latest = records[-1] if isinstance(records, list) else records
    citation_map = latest["citation_map"]
    replacements = {
        label: f"{label}{_format_gbt_reference(enriched(item))}"
        for label, item in citation_map.items()
    }

    shutil.copyfile(SOURCE, OUTPUT)
    doc = Document(OUTPUT)
    for paragraph in doc.paragraphs:
        if "发动机工程师为了安全起见" in paragraph.text and "[10]" in paragraph.text:
            text = paragraph.text.replace("[10]", "[5]")
            paragraph._p.clear_content()
            paragraph.add_run(text)
        text = paragraph.text.strip()
        if not text.startswith("["):
            continue
        label = text.split("]", 1)[0] + "]"
        if label == "[10]":
            paragraph._element.getparent().remove(paragraph._element)
            continue
        if label not in replacements:
            continue
        paragraph._p.clear_content()
        paragraph.add_run(replacements[label])
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
