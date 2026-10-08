from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path('outputs/商用航空发动机情报系统_线上演示测试方案.docx')
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.68)
sec.left_margin = sec.right_margin = Inches(.75)
for grid in sec._sectPr.findall(qn('w:docGrid')): sec._sectPr.remove(grid)
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
for name, size in [('Normal',11),('Title',22),('Heading 1',16),('Heading 2',12)]:
    st=doc.styles[name]; st.font.name='Microsoft YaHei'; st.font.size=Pt(size); st.font.color.rgb=RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'微软雅黑')
    st.paragraph_format.space_after=Pt(5)
    st.paragraph_format.line_spacing=Pt({'Normal':16,'Title':29,'Heading 1':22,'Heading 2':18}[name])
    st.paragraph_format.space_before=Pt(9 if name.startswith('Heading') else 0)
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); st.element.get_or_add_pPr().append(snap)
doc.styles['Normal'].paragraph_format.widow_control=True
md=[]

def p(text, bold=False):
    para=doc.add_paragraph(); para.add_run(text).bold=bold; md.append(text+'\n'); return para
def h(text,level=1):
    doc.add_heading(text,level=level); md.append('#'*(level+1)+' '+text+'\n')
def page(title):
    h(title); doc.paragraphs[-1].paragraph_format.page_break_before=True
def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Inches(w)
    for i,row in enumerate([headers]+rows):
        cells=t.rows[0].cells if i==0 else t.add_row().cells
        tr=t.rows[i]._tr.get_or_add_trPr(); tr.append(OxmlElement('w:cantSplit'))
        if i==0: tr.append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(row):
            cell=cells[j]; cell.width=Inches(widths[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props=cell._tc.get_or_add_tcPr()
            shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'E9EFF5' if i==0 else 'FFFFFF'); props.append(shading)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:w'),'75'); el.set(qn('w:type'),'dxa'); margins.append(el)
            props.append(margins)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                el=OxmlElement('w:'+side); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'D9D9D9'); borders.append(el)
            props.append(borders)
            para=cell.paragraphs[0]; para.paragraph_format.space_before=Pt(1); para.paragraph_format.space_after=Pt(2); para.paragraph_format.line_spacing=Pt(14)
            if i==0: para.paragraph_format.keep_with_next=True
            run=para.add_run(str(value)); run.font.size=Pt(10); run.bold=i==0
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    md.append('| '+' | '.join(headers)+' |\n| '+' | '.join(['---']*len(headers))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in rows)+'\n')
def case(title,steps,expected,show):
    h(title,2)
    p('操作：'+steps)
    p('预期与判定：'+expected)
    p('展示与留证：'+show)

doc.add_paragraph('商用航空发动机情报系统\n线上演示测试方案',style='Title')
md.append('# 商用航空发动机情报系统线上演示测试方案\n')
p('演示日期：2026年9月30日　　建议时长：20分钟　　版本：V1.0')
h('1 演示目标与测试场景')
p('以普惠GTF技术问题与维修保障跟踪为例，展示从研究需求输入、资料检索、报告生成到来源追溯和成果导出的完整流程。现场启动真实任务，生成期间讲解资料库与预生成样例，任务完成后切回本次结果。')
p('本方案用于演示准备和功能测试。测试结果由彩排及现场执行后填写；指标测评可作为辅助展示，当前不能据此宣称准确率已经达标。')
h('建议固定的任务配置',2)
table(['配置项','演示设置'],[
 ['入口','http://127.0.0.1:8000；使用原有浏览器和用户配置'],
 ['任务模板与领域需求模型','不使用模板；领域需求模型选择自动匹配'],
 ['报告类型与长短','技术动态研报；短报告（DeepSeek Flash）'],
 ['生成大模型','DeepSeek Flash；确认页面提示已配置'],
 ['检索范围','论文库、专利池、用户资料、Web；与彩排保持一致'],
 ['精读及优先网址','精读/抓取上限为5；保留彩排验证过的优先网址设置'],
],[1.65,5.35])
h('可直接复制的测试任务',2)
p('GTF发动机技术问题和市场影响跟踪，重点关注PW1000G、PW1100G、PW1500G、PW1900G、粉末金属污染、受影响部件、适航响应、维修保障能力、交付延误和市场影响。请区分型号、时间及计划与实际状态，结合本地资料和公开网页形成中文研究报告，保留可追溯的引用。')
p('执行原则：演示与彩排采用同一段任务描述、同一模型和同一检索设置。现场只启动一项研究任务，避免重复提交。')

page('2 演示前准备与完整彩排')
h('今晚完成的准备',2)
table(['检查项','操作与完成标准'],[
 ['运行环境','在项目根目录运行start_system.bat；页面可以打开，“状态”中连接正常。后台窗口保持运行。'],
 ['模型与网络','确认生成模型可选；用一次完整彩排验证搜索、模型调用及网页抓取。仅连接正常不能证明外部接口可用。'],
 ['资料库','在“论文池概览”搜索“2025年民用航空动力进展”，确认能找到资料，并准备对应PDF原文。'],
 ['上传功能','如需展示，今晚先用已授权、可复制文字的小型PDF彩排；确认文件入库及索引状态。主任务启动前完成。'],
 ['会议共享','测试浏览器与Word/PDF窗口切换；确认观众能看清正文。关闭通知，不共享含密钥的配置文件。'],
 ['备用材料','提前打开af0e主报告、PDF、指标附件及原文；保存本次彩排的关键截图。备用材料标明“预生成样例”。'],
],[1.25,5.75])
h('完整彩排步骤',2)
p('1. 按第1页配置启动一次任务，记录点击启动、报告可读、文件导出完成的时间。')
p('2. 按第4至5页执行T01至T08，完成报告阅读、引用核对和文件打开。不要只检查“生成成功”的提示。')
p('3. 留存配置截图、运行阶段截图、报告首页、来源对照和下载结果；记录任务编号与输出文件名。')
p('4. 固定彩排通过的配置和文件位置。演示前避免升级依赖、重建全库索引或临时切换未验证模型。')
h('备用材料的位置',2)
p('项目目录：D:\\商发多智能体情报收集\\commercial-aircraft-engine-intelligence')
p('在outputs中搜索“af0e”，可找到主报告DOCX、PDF、Markdown及指标附件。原文位于local_docs\\all_papers_pool\\2025年民用航空动力进展_廖忠权.pdf。')
p('彩排通过条件：能完成任务配置、看到实际运行进展、打开报告及导出文件，并完成至少一条本地来源和一条网页来源的人工对照。若生成链路仍阻断，现场明确采用预生成结果演示。')

page('3 现场展示顺序与时间安排')
p('采用20分钟安排，主任务应在第3分钟以前启动。后台生成期间展示独立打开的文件，不在运行页面加载历史报告，以免替换正在展示的任务内容。')
table(['时间','操作步骤','重点展示内容'],[
 ['0—1分钟','介绍任务和工作台','研究需求、资料来源、可交付成果'],
 ['1—3分钟','填入任务并点击“启动情报任务”','模型、报告形式、检索范围及真实启动状态'],
 ['3—5分钟','观察“运行日志”和“状态”','任务编号、当前阶段、已用时间和检索进展'],
 ['5—8分钟','展示资料库及事先打开的原文','本地论文、专利、用户资料的组织及读取方式'],
 ['8—12分钟','切到预生成报告讲解结构','摘要、技术分析、维修保障影响、图表及结论'],
 ['12—15分钟','展示样例的引用对照','正文判断、参考文献、具体原文段落或网页'],
 ['15—18分钟','切回现场任务；完成后打开结果','本次报告内容、Word/PDF下载和质量提示'],
 ['18—20分钟','简述测评、记录结果并答疑','哪些流程已完成、哪些结论仍需人工复核'],
],[1.0,2.35,3.65])
h('到点后的处理',2)
p('第15分钟若现场任务仍在运行，继续展示预生成报告的导出和来源核对；保留后台任务，不反复点击启动。第18分钟统一进入结果说明与答疑，不让等待占满演示时间。')
p('af0e记录的报告生成耗时为11.70分钟，可用作彩排参考。新任务耗时受网络、资料量和模型响应影响，现场不承诺复现相同用时。')
h('只有10分钟时的精简安排',2)
p('0—2分钟：配置并启动；2—4分钟：日志与资料库；4—7分钟：预生成报告和一条来源对照；7—9分钟：下载与质量提示；9—10分钟：说明后台进度并答疑。新任务不以10分钟内完成为前提。')
h('开场参考话术',2)
p('“今天以普惠GTF技术问题和维修保障跟踪为例，展示系统如何把研究需求转成检索任务，结合本地资料和公开网页形成报告，并保留来源供复核。我会现场启动一项任务，同时用一份预生成样例讲解结果结构。”')

page('4 基础功能与任务运行测试')
case('T01 启动与连接检查',
 '运行start_system.bat，打开http://127.0.0.1:8000；点击右上角“状态”，查看连接状态和当前模型；确认生成模型选项可以加载。',
 '首页、模型配置和资料库正常加载；状态无持续连接错误。未通过时先处理服务问题，不启动研究。',
 '工作台总览和状态面板截图；记录启动时间。')
case('T02 本地资料检索与上传',
 '在“论文池概览”搜索“2025年民用航空动力进展”，查看命中的文件和摘要；打开准备好的对应原文。上传为可选项：选择“用户资料库”，选取授权文件，点击“上传并加入检索范围”。',
 '可找到目标资料并读取原文；如执行上传，文件出现在对应资料库且索引状态明确。索引仍在处理时记录为未完成，不表述为全文已可检索。',
 '资料列表、命中文件、原文页；可选展示上传后的排队、索引中或完成状态。')
case('T03 任务配置与提交',
 '复制第1页测试任务，依次核对报告类型、长短、生成模型、检索范围、抓取上限和优先网址；截图后点击一次“启动情报任务”。',
 '系统接受任务并进入运行；日志或状态出现当前任务及阶段。配置必须与彩排一致。',
 '强调用户能控制研究对象、资料范围和报告形式；保存配置截图、实际任务编号及启动时间。')
case('T04 运行过程可见性',
 '观察“运行日志”和“状态”，分别在需求分析、检索研究、汇总写作或后续阶段保存截图；讲解时以实际出现的阶段名称为准。',
 '阶段或日志持续更新；若出现错误，能够识别出错环节。页面停留在同一阶段不直接判为失败，结合最后活动时间判断。',
 '展示任务分工和处理过程。检索覆盖可视化只说明界面展示的覆盖信息，不把进度条当作事实准确率。')

page('5 报告成果与来源追溯测试')
case('T05 报告内容检查',
 '任务完成后展开“报告预览”；依次查看标题、摘要、技术问题、适航响应、维修保障、市场影响、结论和参考文献。',
 '正文可读、章节连续、没有失败占位文本；图表如出现，检查图题和来源。未生成图表时如实记录，不单独视为失败。',
 '挑选一段分析说明“事实如何支持判断”，避免逐页朗读。记录缺失章节或明显内容问题。')
case('T06 正文与原文对照',
 '本地示例：af0e市场影响段中的“2025年PW1000G交付1055台”对应参考文献[6]；打开《2025年民用航空动力进展》PDF第6页并搜索“1055”。网页示例：核对af0e参考文献[10]的腾讯新闻链接与正文停场率表述。',
 '实体、数值、年份、统计对象与来源内容一致；网页能打开只是第一步，还应找到支持该判断的具体内容。现场新报告编号可能改变，须以当次参考文献为准。',
 '按“正文一句话—参考文献—原文段落”展示。任一项对不上就标为待复核，记录原因。')
case('T07 导出与历史查看',
 '点击Word、PDF下载按钮，分别打开文件并检查首页、正文、参考文献；任务完成后点击“历史”，搜索本次任务并打开对应成果。',
 '文件存在且可打开，内容与本次任务一致；历史链接指向正确报告。历史保存在当前浏览器，换浏览器或清空数据后不保证存在。',
 '展示可编辑的Word、便于传阅的PDF及历史访问；记录文件名和导出完成时间。')
case('T08 测评附件与可选问答',
 '打开“指标测评”，查看关联任务、耗时、样本数和待复核状态。若新任务完成后出现“报告问答”，可问：“请依据当前报告概括影响维修周转的三个因素，并指出对应章节。”',
 '附件关联正确，不把待核验计作已通过。问答仅作可选展示；历史报告加载后问答入口隐藏，不安排对历史报告继续提问。',
 '展示质量追踪机制；如执行问答，检查回答是否对应本报告。不要把回答当作新增事实的独立证据。')

page('6 异常处理与现场说明')
table(['情况','现场处理','建议说明'],[
 ['页面或连接异常','检查后台是否运行；不反复启动多个服务。约1分钟内不能恢复时，切换预生成文件。','当前演示环境连接异常，先展示已生成成果及来源对照。'],
 ['外部接口慢或报错','保留任务编号和错误截图；继续讲解资料库及样例，不连续重复提交。','这一步依赖外部服务，当前响应较慢，后台任务状态可以继续查看。'],
 ['报告生成时间较长','第15分钟切换备用内容，第18分钟进入答疑；现场任务继续运行。','当前任务仍在处理，接下来用预生成样例说明成果形式。'],
 ['网页打不开','打开已提前保存并标记日期的原文截图或摘录；记录实时访问失败。','这里展示的是预先保存的原文，当前网页访问情况另行记录。'],
 ['Word或PDF导出失败','先展示可用的报告预览或另一种格式；保留失败状态和文件名。','正文已生成，但这一格式导出异常，本项记为未通过。'],
 ['测评未全部完成','展示实际状态与待复核数量，简述后回到业务流程。','当前测评用于辅助定位问题，仍有待核验项，尚不能作完整准确率结论。'],
],[1.3,2.85,2.85])
h('常见提问与回答口径',2)
p('问：报告来自哪里？\n答：本地论文、专利、用户资料和公开网页共同参与检索；具体判断应查看正文引用及对应原文，资料被检索到不等于已支持全部结论。')
p('问：怎样体现多智能体协同？\n答：结合当前运行日志，说明需求规划、专题研究、本地资料补充、汇总写作与成稿复查各环节的分工，不承诺所有环节同时运行。')
p('问：准确率是否已经达标？\n答：当前版本仍存在未完成核验的样本，现阶段不作达标承诺。今天重点验证生成、追溯和导出流程，准确率验收另行组织完整测试。')
p('问：为什么展示已有文件？\n答：这些明确标记为预生成样例，用于在现场任务运行期间讲解成果；现场生成的报告会用本次任务编号单独识别。')

page('7 测试记录与演示前核对')
p('执行人：________________　日期：________________　浏览器：________________')
p('任务编号：________________________　配置截图编号：________________________')
p('启动时间：____________　报告可读时间：____________　导出完成时间：____________')
p('展示方式：现场任务 / 预生成样例 / 两者结合　　主报告文件名：__________________')
table(['编号','测试内容','结果','证据或问题'],[
 ['T01','启动与连接','待执行',''],
 ['T02','资料检索；上传可选','待执行',''],
 ['T03','任务配置与提交','待执行',''],
 ['T04','运行阶段与日志','待执行',''],
 ['T05','报告内容与结构','待执行',''],
 ['T06','本地及网页来源对照','待执行',''],
 ['T07','Word/PDF与历史访问','待执行',''],
 ['T08','测评附件；问答可选','待执行',''],
],[.55,2.4,.95,3.1])
p('结果填写“通过、未通过、未完成、未执行”。可选项未执行时注明；不要预填通过。截图建议用“任务编号_步骤编号_内容”命名。')
h('演示开始前30分钟',2)
p('1. 打开系统，确认后台运行、网络连接、模型配置和资料列表正常。')
p('2. 将测试任务复制到剪贴板；准备主报告、PDF、测评附件和PDF原文第6页。')
p('3. 测试会议共享和窗口切换；关闭通知，确认浏览器文字大小适合远程观看。')
p('4. 分清“现场任务”窗口和“预生成样例”窗口；准备计时器和本页记录表。')
h('结束前必须讲清楚',2)
p('说明本次实际完成了哪些操作、采用了哪份报告、哪些项仍待处理。只有现场真实完成的操作计入现场结果；预生成文件的展示和检查单独记录。')
p('演示结果：____________________________________________________________')
p('遗留问题与后续安排：__________________________________________________')

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)
doc.core_properties.title='商用航空发动机情报系统线上演示测试方案'
doc.core_properties.subject='演示准备 操作步骤 展示内容 测试记录'
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
OUT.with_suffix('.md').write_text('\n'.join(md),encoding='utf8')
print(OUT)
