"""Create the two revised DOCX artifacts using the bundled document runtime."""
from pathlib import Path
import re
from urllib.parse import unquote
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image

root = Path.cwd()

def runs(p, text):
    text = re.sub(r'\[\[(\d+)\]\]\(#ref-\d+\)', r'[\1]', text)
    text = re.sub(r'\[([^\]]+)\]\(#.*?\)', r'\1', text).replace('**', '')
    pattern = re.compile(r'\[([^\]]+)\]\((https?://[^)]+)\)')
    start = 0
    for match in pattern.finditer(text):
        p.add_run(text[start:match.start()])
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(match[2], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
        run = OxmlElement('w:r'); t = OxmlElement('w:t'); t.text = match[1]; run.append(t); link.append(run); p._p.append(link)
        start = match.end()
    p.add_run(text[start:])

def build(path):
    doc = Document()
    for borders in doc.styles.element.xpath('.//w:pBdr'):
        borders.getparent().remove(borders)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ('Normal','Title','Heading 1','Heading 2','Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')
        style.font.color.rgb = RGBColor(0,0,0)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.35
    doc.styles['Normal'].paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(19)
    lines = path.read_text(encoding='utf8').splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip(); index += 1
        if not line or line.startswith('<a ') or line == '---': continue
        if line.startswith('|'):
            rows = [line]
            while index < len(lines) and lines[index].strip().startswith('|'):
                rows.append(lines[index].strip()); index += 1
            rows = [row for row in rows if not re.fullmatch(r'[| :\-]+', row)]
            cells = [[c.strip() for c in row.strip('|').split('|')] for row in rows]
            table = doc.add_table(rows=0, cols=max(map(len,cells))); table.style = 'Table Grid'
            header = cells[0]
            weights = ([.09,.16,.12,.63] if header == ['编号','对象','判定','理由'] else
                       [.14,.18,.68] if header[-1] == '原因' else
                       [.14,.48,.38] if header == ['引用','来源','结果'] else
                       [.15,.12,.18,.55] if header == ['服务','状态','HTTP状态','说明'] else
                       [.33,.20,.29,.18] if header[0] == '指标' else
                       [1/len(header)] * len(header))
            table.autofit = False
            for col, weight in zip(table.columns, weights): col.width = Inches(6.77*weight)
            for row_index, row in enumerate(cells):
                tcells = table.add_row().cells
                no_split=OxmlElement('w:cantSplit'); table.rows[-1]._tr.get_or_add_trPr().append(no_split)
                for col, value in enumerate(row):
                    tcells[col].width = Inches(6.77*weights[col])
                    runs(tcells[col].paragraphs[0], value)
                    fmt=tcells[col].paragraphs[0].paragraph_format
                    fmt.line_spacing=1.1; fmt.space_after=Pt(4); fmt.space_before=Pt(4)
                    if row_index == 0: fmt.keep_with_next = True
                    for r in tcells[col].paragraphs[0].runs: r.font.size=Pt(9); r.bold=row_index==0
                if row_index==0:
                    repeat=OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
        elif line.startswith('!['):
            match = re.match(r'!\[.*?\]\((.*?)\)', line)
            if match:
                img = unquote(match[1]); target = root / img.lstrip('/') if img.startswith('/outputs/') else Path(img)
                if target.is_file():
                    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    with Image.open(target) as im: width,height=im.size
                    scale=min(6.5/width, 6/height)
                    p.add_run().add_picture(str(target),width=Inches(width*scale),height=Inches(height*scale))
        elif line.startswith('#'):
            level=len(line)-len(line.lstrip('#')); title=line[level:].strip()
            p=doc.add_paragraph(style='Title' if level==1 else 'Heading '+str(min(level-1,3)))
            runs(p,title)
        else:
            p=doc.add_paragraph(); runs(p,line)
            if re.match(r'^\[\d+\]', line):
                p.paragraph_format.keep_together = True
                p.paragraph_format.line_spacing = 1.1
                p.paragraph_format.space_after = Pt(5)
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple'); f.set(qn('w:instr'),'PAGE'); footer._p.append(f)
    destination=path.with_suffix('.docx'); doc.save(destination); print(destination)

for path in sorted(Path('outputs').glob('*3465_修订版*.md')): build(path)
