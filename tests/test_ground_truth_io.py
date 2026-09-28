import json
import math
import zipfile
from io import BytesIO

import pytest
from openpyxl import Workbook

from gpt_researcher.evaluation import entity_evaluator, ground_truth_io
from gpt_researcher.evaluation.ground_truth_io import (
    ALLOWED_SUFFIXES,
    MAX_UPLOAD_BYTES,
    GroundTruthValidationError,
    ground_truth_path_for_task,
    load_ground_truth_upload,
    persist_ground_truth_upload,
)


TASK = "engine-reliability"


def _valid_payload(**overrides):
    payload = {
        "task": TASK,
        "entities": [
            {
                "type": "参数",
                "name": "推力",
                "aliases": ["thrust"],
                "value": 100,
                "unit": "kN",
            }
        ],
    }
    payload.update(overrides)
    return payload


def _write_json(tmp_path, payload, name="truth.json"):
    path = tmp_path / name
    path.write_text(json.dumps(payload, ensure_ascii=False, allow_nan=True), encoding="utf-8")
    return path


def _xlsx_bytes(rows):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "实体表"
    for row in rows:
        worksheet.append(row)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _write_xlsx(tmp_path, rows, name="truth.xlsx"):
    path = tmp_path / name
    path.write_bytes(_xlsx_bytes(rows))
    return path


def _assert_error(exc_info, code, message_part=""):
    error = exc_info.value
    assert error.code == code
    assert error.message == str(error)
    if message_part:
        assert message_part in error.message


def test_json_and_xlsx_normalize_to_the_same_canonical_entities(tmp_path):
    json_path = _write_json(tmp_path, _valid_payload())
    xlsx_path = _write_xlsx(
        tmp_path,
        [
            ["类别", "名称", "别名", "数值", "单位", "容差"],
            ["参数", "推力", "thrust", 100, "kN", None],
        ],
    )

    json_truth = load_ground_truth_upload(json_path, TASK)
    xlsx_truth = load_ground_truth_upload(xlsx_path, TASK)

    assert ALLOWED_SUFFIXES == {".json", ".xlsx"}
    assert json_truth == xlsx_truth == {
        "task": TASK,
        "entities": [
            {
                "type": "parameter",
                "name": "推力",
                "aliases": ["thrust"],
                "value": 100,
                "unit": "kN",
                "tolerance": 0.01,
            }
        ],
    }


@pytest.mark.parametrize(
    ("name", "content", "code"),
    [
        pytest.param("truth.txt", b"{}", "unsupported_file_type", id="unsupported"),
        pytest.param("truth.json", b"", "empty_file", id="empty"),
        pytest.param("truth.json", b"{" + b" " * MAX_UPLOAD_BYTES, "file_too_large", id="oversized"),
    ],
)
def test_upload_rejects_unsupported_empty_and_oversized_files(tmp_path, name, content, code):
    path = tmp_path / name
    path.write_bytes(content)

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, code)


@pytest.mark.parametrize(
    ("payload", "code"),
    [
        ([{"type": "model", "name": "LEAP"}], "invalid_schema"),
        (_valid_payload(task="another-task"), "task_mismatch"),
        (_valid_payload(entities=[]), "empty_entities"),
        (_valid_payload(entities=[{"name": "LEAP"}]), "invalid_entity"),
        (_valid_payload(entities=[{"type": "model"}]), "invalid_entity"),
        (_valid_payload(entities=[{"type": "model", "name": "LEAP", "aliases": "bad"}]), "invalid_aliases"),
        (_valid_payload(entities=[{"type": "parameter", "name": "ratio", "value": True}]), "invalid_parameter"),
        (_valid_payload(entities=[{"type": "parameter", "name": "ratio", "value": float("nan")}]), "invalid_parameter"),
        (_valid_payload(entities=[{"type": "parameter", "name": "ratio", "tolerance": -0.01}]), "invalid_parameter"),
        (_valid_payload(entities=[{"type": "parameter", "name": "ratio", "tolerance": math.inf}]), "invalid_parameter"),
    ],
)
def test_json_schema_validation_has_stable_error_codes(tmp_path, payload, code):
    path = _write_json(tmp_path, payload)

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, code)


def test_missing_task_is_injected_and_duplicate_normalized_entities_are_rejected(tmp_path):
    missing_task_payload = _valid_payload()
    del missing_task_payload["task"]
    injected_path = _write_json(tmp_path, missing_task_payload)
    assert load_ground_truth_upload(injected_path, TASK)["task"] == TASK

    duplicate_path = _write_json(
        tmp_path,
        _valid_payload(
            entities=[
                {"type": "型号", "name": "LEAP-1A"},
                {"type": "model", "name": "leap 1a"},
            ]
        ),
        "duplicate.json",
    )
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(duplicate_path, TASK)

    _assert_error(exc_info, "duplicate_entity")


def test_xlsx_requires_columns_and_reports_sheet_name_and_actual_row(tmp_path):
    missing_columns = _write_xlsx(tmp_path, [["类别", "别名"], ["型号", "LEAP"]])
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(missing_columns, TASK)
    _assert_error(exc_info, "missing_columns", "实体表")

    invalid_row = _write_xlsx(
        tmp_path,
        [["类别", "名称"], ["型号", "LEAP"], ["参数", None]],
        "invalid-row.xlsx",
    )
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(invalid_row, TASK)
    _assert_error(exc_info, "invalid_entity", "实体表，第 3 行")


def test_xlsx_formula_without_cached_value_is_rejected_with_sheet_and_row(tmp_path):
    path = _write_xlsx(
        tmp_path,
        [["类别", "名称", "数值"], ["参数", "推力", "=100"]],
    )

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "formula_without_cached_value", "实体表，第 2 行")


def test_persist_upload_uses_hashed_canonical_json_inside_destination(tmp_path):
    destination = tmp_path / "stored"
    metadata = persist_ground_truth_upload(
        json.dumps(_valid_payload(), ensure_ascii=False).encode("utf-8"),
        "../../outside.json",
        TASK,
        destination,
    )
    target = ground_truth_path_for_task(TASK, destination)

    assert target.parent == destination
    assert target.exists()
    assert metadata["stored_name"] == target.name
    assert len(metadata["sha256"]) == 64
    assert metadata["entity_count"] == 1
    assert metadata["category_counts"] == {"parameter": 1}
    assert json.loads(target.read_text(encoding="utf-8"))["entities"][0]["tolerance"] == 0.01
    assert not (tmp_path / "outside.json").exists()
    assert all(path.parent == destination for path in destination.iterdir())
    assert list(destination.iterdir()) == [target]


def test_parameter_string_value_must_be_a_finite_single_number(tmp_path):
    path = _write_json(
        tmp_path,
        _valid_payload(entities=[{"type": "parameter", "name": "ratio", "value": "garbage"}]),
    )

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "invalid_parameter")


def test_duplicate_keys_preserve_negative_sign_and_normalize_numeric_equivalence(tmp_path):
    distinct = _write_json(
        tmp_path,
        _valid_payload(
            entities=[
                {"type": "parameter", "name": "ratio", "value": -1},
                {"type": "parameter", "name": "ratio", "value": 1},
            ]
        ),
        "distinct.json",
    )
    assert len(load_ground_truth_upload(distinct, TASK)["entities"]) == 2

    equivalent = _write_json(
        tmp_path,
        _valid_payload(
            entities=[
                {"type": "parameter", "name": "ratio", "value": 100},
                {"type": "parameter", "name": "ratio", "value": "100.0"},
            ]
        ),
        "equivalent.json",
    )
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(equivalent, TASK)
    _assert_error(exc_info, "duplicate_entity")


def test_xlsx_all_formula_data_row_without_cached_values_is_not_ignored(tmp_path):
    path = _write_xlsx(
        tmp_path,
        [["类别", "名称", "数值"], ["=\"参数\"", "=\"推力\"", "=100"]],
    )

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "formula_without_cached_value", "实体表，第 2 行")


def test_persist_cleans_a_temporary_file_when_canonical_json_write_fails(tmp_path, monkeypatch):
    destination = tmp_path / "stored"

    def fail_json_dump(*_args, **_kwargs):
        raise OSError("disk full")

    monkeypatch.setattr(ground_truth_io.json, "dump", fail_json_dump)
    with pytest.raises(OSError, match="disk full"):
        persist_ground_truth_upload(
            json.dumps(_valid_payload(), ensure_ascii=False).encode("utf-8"),
            "truth.json",
            TASK,
            destination,
        )

    assert list(destination.iterdir()) == []


def test_persist_cleans_a_temporary_file_when_upload_write_fails(tmp_path, monkeypatch):
    destination = tmp_path / "stored"
    real_named_temporary_file = ground_truth_io.tempfile.NamedTemporaryFile

    class FailingTemporaryFile:
        def __init__(self, *args, **kwargs):
            self._temporary = real_named_temporary_file(*args, **kwargs)
            self.name = self._temporary.name

        def __enter__(self):
            self._temporary.__enter__()
            return self

        def __exit__(self, *args):
            return self._temporary.__exit__(*args)

        def write(self, _content):
            raise OSError("disk full")

    monkeypatch.setattr(ground_truth_io.tempfile, "NamedTemporaryFile", FailingTemporaryFile)
    with pytest.raises(OSError, match="disk full"):
        persist_ground_truth_upload(
            json.dumps(_valid_payload(), ensure_ascii=False).encode("utf-8"),
            "truth.json",
            TASK,
            destination,
        )

    assert list(destination.iterdir()) == []


def test_nonconvertible_tolerance_returns_a_stable_validation_error(tmp_path):
    path = _write_json(
        tmp_path,
        _valid_payload(entities=[{"type": "parameter", "name": "ratio", "tolerance": 10**400}]),
    )

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "invalid_parameter")


@pytest.mark.parametrize("declared_task", [None, ""])
def test_explicit_missing_like_task_values_are_not_injected(tmp_path, declared_task):
    path = _write_json(tmp_path, _valid_payload(task=declared_task))

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "task_mismatch")


def test_explicit_null_aliases_are_invalid(tmp_path):
    path = _write_json(
        tmp_path,
        _valid_payload(entities=[{"type": "model", "name": "LEAP", "aliases": None}]),
    )

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "invalid_aliases")


def test_legacy_evaluator_fallback_keeps_shared_file_size_validation(tmp_path, monkeypatch):
    task = "legacy-oversized"
    path = tmp_path / f"{task}.json"
    path.write_bytes(json.dumps({"entities": [{"name": "LEAP"}]}).encode() + b" " * MAX_UPLOAD_BYTES)
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)

    result = entity_evaluator.load_ground_truth(task)

    assert result["status"] == "invalid_ground_truth"
    assert result["error_code"] == "file_too_large"


def test_legacy_evaluator_fallback_keeps_shared_file_type_validation(tmp_path, monkeypatch):
    path = tmp_path / "legacy.txt"
    path.write_text(json.dumps([{"name": "LEAP"}]), encoding="utf-8")
    monkeypatch.setattr(entity_evaluator, "_ground_truth_candidates", lambda _task: [path])

    result = entity_evaluator.load_ground_truth("legacy-type")

    assert result["status"] == "invalid_ground_truth"
    assert result["error_code"] == "unsupported_file_type"


def test_corrupt_uploads_have_stable_parse_errors(tmp_path):
    invalid_json = tmp_path / "broken.json"
    invalid_json.write_bytes(b"{")
    invalid_xlsx = tmp_path / "broken.xlsx"
    invalid_xlsx.write_bytes(b"not an xlsx")

    for path, code in ((invalid_json, "invalid_json"), (invalid_xlsx, "invalid_excel")):
        with pytest.raises(GroundTruthValidationError) as exc_info:
            load_ground_truth_upload(path, TASK)
        _assert_error(exc_info, code)


def test_xlsx_uses_first_visible_sheet_and_splits_both_semicolon_forms(tmp_path):
    workbook = Workbook()
    hidden = workbook.active
    hidden.title = "隐藏"
    hidden.append(["错误", "表头"])
    hidden.sheet_state = "hidden"
    visible = workbook.create_sheet("标准答案")
    visible.append(["类别", "名称", "别名"])
    visible.append(["型号", "LEAP", "LEAP-1A; LEAP-1B；LEAP"])
    path = tmp_path / "visible.xlsx"
    workbook.save(path)

    canonical = load_ground_truth_upload(path, TASK)

    assert canonical["entities"] == [
        {"type": "model", "name": "LEAP", "aliases": ["leap1a", "leap1b", "leap"]}
    ]


@pytest.mark.parametrize("value", ["1e1000000", "1e309"])
def test_parameter_values_must_be_finite_for_the_existing_matcher(tmp_path, value):
    path = _write_json(
        tmp_path,
        _valid_payload(entities=[{"type": "parameter", "name": "ratio", "value": value}]),
    )

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "invalid_parameter")


def test_duplicate_parameter_keys_include_embedded_units(tmp_path):
    distinct_units = _write_json(
        tmp_path,
        _valid_payload(
            entities=[
                {"type": "parameter", "name": "thrust", "value": "100 kN"},
                {"type": "parameter", "name": "thrust", "value": "100 N"},
            ]
        ),
        "distinct-units.json",
    )
    assert len(load_ground_truth_upload(distinct_units, TASK)["entities"]) == 2

    equivalent_units = _write_json(
        tmp_path,
        _valid_payload(
            entities=[
                {"type": "parameter", "name": "thrust", "value": "100 kN"},
                {"type": "parameter", "name": "thrust", "value": 100, "unit": "kN"},
            ]
        ),
        "equivalent-units.json",
    )
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(equivalent_units, TASK)
    _assert_error(exc_info, "duplicate_entity")


def test_delayed_corrupt_xlsx_sheet_xml_returns_safe_error_and_releases_file(tmp_path):
    path = _write_xlsx(tmp_path, [["类别", "名称"], ["型号", "LEAP"]])
    replacement = tmp_path / "replacement.xlsx"
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                content = (
                    b'<?xml version="1.0" encoding="UTF-8"?>'
                    b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                    b'<dimension ref="A1:B2"/><sheetData><row'
                )
            target.writestr(info, content)
    replacement.replace(path)

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "invalid_excel")
    path.unlink()
    assert not path.exists()


def test_sparse_xlsx_dimension_is_rejected_before_workbook_iteration(tmp_path, monkeypatch):
    path = tmp_path / "sparse.xlsx"
    sheet_xml = (
        b'<?xml version="1.0" encoding="UTF-8"?>'
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<dimension ref="A1:XFD1048576"/><sheetData><row r="1"/></sheetData></worksheet>'
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)

    def must_not_open(*_args, **_kwargs):
        raise AssertionError("workbook iteration must not start")

    monkeypatch.setattr(ground_truth_io, "load_workbook", must_not_open)
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)

    _assert_error(exc_info, "xlsx_too_large")


@pytest.mark.parametrize(
    "content",
    [
        pytest.param(b"[" * 1100 + b"0" + b"]" * 1100, id="deep-array"),
        pytest.param(b"{" + b'"n":' + b"9" * 5000 + b"}", id="huge-integer"),
    ],
)
def test_deep_or_huge_json_numbers_return_safe_error_for_upload_and_evaluator(tmp_path, monkeypatch, content):
    path = tmp_path / "unsafe.json"
    path.write_bytes(content)

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)
    _assert_error(exc_info, "invalid_json")

    evaluator_path = tmp_path / "unsafe-task.json"
    evaluator_path.write_bytes(content)
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)
    result = entity_evaluator.load_ground_truth("unsafe-task")
    assert result["status"] == "invalid_ground_truth"
    assert result["error_code"] == "invalid_json"


def test_legacy_synonyms_are_merged_into_aliases_for_matching(tmp_path, monkeypatch):
    task = "legacy-synonyms"
    (tmp_path / f"{task}.json").write_text(
        json.dumps({"entities": [{"name": "LEAP-1A", "category": "model", "synonyms": ["LEAP"]}]}),
        encoding="utf-8",
    )
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)

    truth = entity_evaluator.load_ground_truth(task)
    metrics = entity_evaluator.evaluate_entities_against_ground_truth(
        [{"name": "LEAP", "category": "model"}], truth["entities"]
    )

    assert truth["status"] == "loaded"
    assert truth["entities"][0]["aliases"] == ["leap"]
    assert metrics["overall"]["f1"] == 1.0


def test_legacy_category_overrides_type_and_normalizes_equivalent_task(tmp_path, monkeypatch):
    task = "t"
    (tmp_path / f"{task}.json").write_text(
        json.dumps(
            {"task": "T", "entities": [{"name": "LEAP", "category": "model", "type": "organization"}]}
        ),
        encoding="utf-8",
    )
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)

    truth = entity_evaluator.load_ground_truth(task)

    assert truth["status"] == "loaded"
    assert truth["entities"][0]["category"] == "model"


def test_xlsx_preflight_rejects_excessive_actual_cell_nodes(tmp_path, monkeypatch):
    path = tmp_path / "many-cells.xlsx"
    cells = b"<c r=\"A1\"/>" * 120_000
    sheet_xml = (
        b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<dimension ref="A1:B2"/><sheetData><row r="1">' + cells + b"</row></sheetData></worksheet>"
    )
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/worksheets/sheet1.xml", sheet_xml)

    monkeypatch.setattr(ground_truth_io, "load_workbook", lambda *_args, **_kwargs: pytest.fail("opened"))
    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)
    _assert_error(exc_info, "xlsx_too_large")


def test_evaluator_loads_canonical_hashed_ground_truth_after_persist(tmp_path, monkeypatch):
    metadata = persist_ground_truth_upload(
        json.dumps(_valid_payload(), ensure_ascii=False).encode("utf-8"), "truth.json", TASK, tmp_path
    )
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)

    result = entity_evaluator.load_ground_truth(TASK)

    assert result["status"] == "loaded"
    assert result["path"].endswith(metadata["stored_name"])


def test_evaluator_rejects_oversized_file_before_json_parsing(tmp_path, monkeypatch):
    task = "oversized-evaluator"
    (tmp_path / f"{task}.json").write_bytes(b"{" + b" " * (MAX_UPLOAD_BYTES + 1))
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)
    monkeypatch.setattr(ground_truth_io.json, "load", lambda _handle: pytest.fail("parsed oversized file"))

    result = entity_evaluator.load_ground_truth(task)

    assert result["status"] == "invalid_ground_truth"
    assert result["error_code"] == "file_too_large"


@pytest.mark.parametrize("declared_task", [None, ""])
def test_legacy_evaluator_maps_null_or_empty_task_to_requested_task(tmp_path, monkeypatch, declared_task):
    task = "legacy-no-task"
    (tmp_path / f"{task}.json").write_text(
        json.dumps({"task": declared_task, "entities": [{"name": "LEAP"}]}), encoding="utf-8"
    )
    monkeypatch.setattr(entity_evaluator, "get_ground_truth_dir", lambda: tmp_path)

    result = entity_evaluator.load_ground_truth(task)

    assert result["status"] == "loaded"
    assert result["entities"][0]["name"] == "LEAP"


def test_unpaired_surrogates_are_rejected_before_persistence(tmp_path):
    content = b'{"entities":[{"type":"model","name":"\\ud800"}]}'
    path = tmp_path / "surrogate.json"
    path.write_bytes(content)

    with pytest.raises(GroundTruthValidationError) as exc_info:
        load_ground_truth_upload(path, TASK)
    _assert_error(exc_info, "invalid_text")

    destination = tmp_path / "stored"
    with pytest.raises(GroundTruthValidationError) as exc_info:
        persist_ground_truth_upload(content, "surrogate.json", TASK, destination)
    _assert_error(exc_info, "invalid_text")
    assert list(destination.iterdir()) == []


def test_persist_rejects_unpaired_surrogate_task_before_hashing(tmp_path):
    destination = tmp_path / "stored"
    content = json.dumps(_valid_payload(), ensure_ascii=False).encode("utf-8")

    with pytest.raises(GroundTruthValidationError) as exc_info:
        persist_ground_truth_upload(content, "truth.json", "\ud800", destination)

    _assert_error(exc_info, "invalid_text")
    assert list(destination.iterdir()) == []
