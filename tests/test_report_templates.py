import asyncio
from io import BytesIO
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from fastapi import FastAPI
from fastapi.testclient import TestClient

from backend.reporting import template_store as store
from backend.reporting.document_export import render_word
from backend.reporting.template_api import PREVIEW_REPORT, router


class ReportTemplateTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        patcher = patch.object(store, 'TEMPLATE_DIR', self.root / 'templates')
        patcher.start()
        self.addCleanup(patcher.stop)
        environment = patch.dict(os.environ, {'REPORT_EXPORT_UPDATE_WORD_FIELDS': '0'})
        environment.start()
        self.addCleanup(environment.stop)

    def fixture(self):
        doc = Document(BytesIO(store.standard_template_bytes()))
        doc.sections[0].left_margin = Mm(32)
        doc.sections[0].right_margin = Mm(18)
        doc.sections[0].header.paragraphs[0].add_run(' 客户页眉')
        doc.sections[0].first_page_header.paragraphs[0].text = '客户首页页眉'
        p = next(p for p in doc.paragraphs if p.text.startswith('此段'))
        p.paragraph_format.line_spacing = Pt(24)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        p.paragraph_format.first_line_indent = Pt(28)
        p.runs[0].font.size = Pt(14)
        body = BytesIO()
        doc.save(body)
        return body.getvalue()

    def imported(self):
        return store.import_template('客户模板.docx', self.fixture())

    def profile(self, info, **extra):
        return store.resolve_template_profile({'template': {'id': info['id'], 'overrides': []}, **extra})

    def test_import_reads_direct_formatting_and_does_not_adopt_chapters_by_default(self):
        info = self.imported()
        profile = self.profile(info)
        self.assertEqual(profile['size']['body'], 14)
        self.assertEqual(profile['layout']['line_spacing_rule'], 'exact')
        self.assertEqual(profile['layout']['line_spacing'], 24)
        self.assertAlmostEqual(profile['layout']['margin_left_mm'], 32, delta=.1)
        self.assertEqual(profile['chapters'], [])
        self.assertIn('技术分析', [item['title'] for item in info['chapters']])
        self.assertEqual(len(store.list_templates()), 1)

    def test_export_keeps_cover_headers_and_replaces_example_body(self):
        info = self.imported()
        profile = self.profile(info)
        path = self.root / 'report.docx'
        render_word(PREVIEW_REPORT, path, format_profile=profile)
        doc = Document(path)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('报告模板试排', text)
        self.assertNotIn('此段为正文排版示例', text)
        self.assertNotIn('{{REPORT_', text)
        self.assertIn('客户页眉', doc.sections[0].header.paragraphs[0].text)
        self.assertEqual(doc.sections[0].first_page_header.paragraphs[0].text, '客户首页页眉')
        self.assertAlmostEqual(doc.sections[0].left_margin.mm, 32, delta=.1)
        self.assertEqual(doc.styles['Normal'].font.size.pt, 14)
        self.assertEqual(doc.styles['Normal'].paragraph_format.line_spacing.pt, 24)
        self.assertEqual(doc.styles['Normal'].paragraph_format.first_line_indent.pt, 28)
        self.assertTrue(doc.styles['Heading 1'].font.bold)

    def test_only_explicit_edits_override_template(self):
        info = self.imported()
        profile = store.resolve_template_profile({
            'template': {'id': info['id'], 'overrides': ['size.body', 'layout.margin_mm', 'layout.cover']},
            'size': {'body': 12}, 'layout': {'margin_mm': 30, 'cover': False},
            'chapters': ['用户章节'],
        })
        self.assertEqual(profile['layout']['line_spacing'], 24)
        path = self.root / 'edited.docx'
        render_word(PREVIEW_REPORT, path, format_profile=profile)
        doc = Document(path)
        self.assertEqual(doc.styles['Normal'].font.size.pt, 12)
        self.assertAlmostEqual(doc.sections[0].left_margin.mm, 30, delta=.1)
        self.assertEqual(doc.paragraphs[0].text, '报告模板试排')
        self.assertEqual(profile['chapters'][0]['title'], '用户章节')

    def test_template_logo_relationship_is_preserved(self):
        doc = Document(BytesIO(self.fixture()))
        image = Path(__file__).resolve().parents[1] / 'frontend' / 'static' / 'aecc-cae-logo.png'
        if not image.is_file():
            image = Path(__file__).resolve().parents[1] / 'static' / 'aecc-cae-logo.png'
        self.assertTrue(image.is_file(), image)
        doc.sections[0].header.paragraphs[0].add_run().add_picture(str(image), width=Mm(20))
        buffer = BytesIO(); doc.save(buffer)
        info = store.import_template('带Logo.docx', buffer.getvalue())
        path = self.root / 'logo.docx'
        render_word(PREVIEW_REPORT, path, format_profile=self.profile(info))
        exported = Document(path)
        self.assertTrue(exported.sections[0].header._element.xpath('.//a:blip'))
        with ZipFile(path) as archive:
            self.assertTrue(any(name.startswith('word/media/') for name in archive.namelist()))

    def test_bad_or_missing_templates_and_multi_section_fail_clearly(self):
        for name, data in [('bad.pdf', b'pdf'), ('bad.docx', b'not a zip')]:
            with self.assertRaises(store.ReportTemplateError):
                store.import_template(name, data)
        with self.assertRaises(store.ReportTemplateError):
            store.resolve_template_profile({'template': {'id': '../file'}})
        doc = Document(); doc.add_section()
        buffer = BytesIO(); doc.save(buffer)
        with self.assertRaisesRegex(store.ReportTemplateError, '单节'):
            store.import_template('多节.docx', buffer.getvalue())

    def test_unrecognized_override_paths_are_ignored(self):
        info = self.imported()
        profile = store.resolve_template_profile({'template': {'id': info['id'], 'overrides': ['bad', 'size.unknown']}})
        self.assertEqual(profile['template']['overrides'], [])
        render_word(PREVIEW_REPORT, self.root / 'safe.docx', format_profile=profile)

    def test_api_upload_catalog_standard_and_preview(self):
        app = FastAPI(); app.include_router(router)
        with TestClient(app) as client:
            self.assertEqual(client.get('/api/report-templates/standard').status_code, 200)
            response = client.post('/api/report-templates', files={'file': ('客户.docx', self.fixture())})
            self.assertEqual(response.status_code, 200, response.text)
            info = response.json()['template']
            self.assertEqual(len(client.get('/api/report-templates').json()['templates']), 1)
            response = client.post('/api/report-templates/preview', json={
                'report_format_profile': {'template': {'id': info['id'], 'overrides': []}}})
            self.assertEqual(response.status_code, 200, response.text[:200])
            self.assertIn('报告模板试排', '\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs))
            self.assertEqual(client.post('/api/report-templates/preview', json={}).status_code, 400)

    def test_template_pdf_converts_generated_word(self):
        from backend.utils import write_md_to_pdf, write_md_to_word
        info = self.imported()
        old = Path.cwd()
        try:
            os.chdir(self.root)
            with patch('backend.utils.convert_word_to_pdf') as convert:
                profile = self.profile(info)
                asyncio.run(write_md_to_word(PREVIEW_REPORT, 'sample', profile))
                asyncio.run(write_md_to_pdf(PREVIEW_REPORT, 'sample', profile))
                self.assertTrue(convert.call_args.args[0].is_file())
                self.assertEqual(convert.call_args.args[0].suffix, '.docx')
        finally:
            os.chdir(old)

    def test_custom_chapters_keep_corresponding_text_without_positional_renaming(self):
        from backend.reporting.formal_report import prepare_formal_report
        report = '# 技术报告\n\n## 摘要\n摘要内容。\n\n## 技术分析\n叶片分析内容。\n\n## 运营影响\n维修排程内容。'
        prepared = prepare_formal_report(report, '航空发动机技术分析', metadata={
            'report_format_profile': {'chapters': ['运营影响', '技术分析', '不存在的章节']}})
        text = prepared.markdown
        self.assertIn('## 1 运营影响\n\n维修排程内容。', text)
        self.assertIn('## 2 技术分析\n\n叶片分析内容。', text)
        self.assertNotIn('## 3 不存在的章节', text)
        self.assertTrue(any('不存在的章节' in warning for warning in prepared.quality['warnings']))


if __name__ == '__main__':
    unittest.main()
