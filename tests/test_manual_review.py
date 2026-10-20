import copy
from io import BytesIO
import json
import subprocess
import sys
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from urllib.parse import quote

from openpyxl import load_workbook

from gpt_researcher.evaluation.manual_review import (
    create_review_package, import_review, score_rows, summarize_runs, ManualReviewError,
)
from gpt_researcher.evaluation.records import EvaluationRecordStore


class ManualReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        (self.root / 'outputs').mkdir()
        self.report = '# 测试报告\n\n## 摘要\nGTF维修120台[1]。\n\n## 1 分析\nPW1100G维修200台[1]。\n\n## 参考文献\n[1]网页。'
        path = self.root / 'outputs/测试.md'
        path.write_text(self.report, encoding='utf-8')
        self.record = {'run_id': 'test-001', 'task': 'GTF', 'report_title': '测试报告',
                       'md_path': quote('outputs/测试.md'), 'published_markdown': self.report,
                       'total_duration_seconds': 300, 'export_status': {'markdown': True}, 'status': 'ready',
                       'citation_map': {'[1]': {'url': 'https://example.test/a'}},
                       'original_source_catalog': {'sources': [{'locator': 'https://example.test/a',
                           'pages': [{'page': 2, 'text': 'GTF维修120台。PW1100G维修200台。'}]}]}}
        self.package = create_review_package(self.record, self.root)

    def workbook(self):
        return load_workbook(self.package['workbook'])

    def fill(self, wb, decision='正确', confirm=True):
        for row in wb['任务与规则'].iter_rows(min_row=2):
            values = {'核验人': '测试人员', '核验日期': '2026-10-19', '全量检查确认': '确认' if confirm else ''}
            if row[0].value in values:
                row[1].value = values[row[0].value]
        for name in ['实体核验', '溯源核验']:
            for row in wb[name].iter_rows(min_row=2):
                row[9].value = decision
                row[10].value = '逐句核对'
                row[11].value = 'GTF维修120台。PW1100G维修200台。'
                row[12].value = 'https://example.test/a 第2页'
        return wb

    def submit(self, wb, record=None):
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        with patch('backend.reporting.document_export.render_word', side_effect=lambda text, path, **kwargs: Path(path).write_bytes(b'docx')):
            return import_review(record or self.record, stream.getvalue(), self.root)

    def test_package_contains_final_body_abstract_and_saved_sources(self):
        self.assertEqual(self.package['source_count'], 2)
        self.assertEqual(self.package['public_url_count'], 1)
        self.assertEqual(self.package['selection_scope'], 'published_body_only')
        html = Path(self.package['materials']).read_text(encoding='utf-8')
        self.assertIn('GTF维修120台', html)
        self.assertIn('source-0', html)
        self.assertTrue(Path(self.package['workbook']).is_file())

    def test_export_cleanup_keeps_sourced_rows_without_extra_cleanup_files(self):
        self.record['original_source_catalog']['sources'][0]['pages'][0]['text'] = '\ufffd' * 20
        self.package = create_review_package(self.record, self.root)
        folder = Path(self.package['workbook']).parent
        manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
        cleanup = manifest['export_cleanup']
        self.assertGreater(cleanup['removed_excerpts'], 0)
        self.assertEqual(len(manifest['rows']), 4)
        self.assertTrue(all(not row['evidence'] for row in manifest['rows']))
        self.assertFalse((folder / 'export_cleanup.json').exists())
        self.assertNotIn('events', cleanup)
        self.assertEqual({p.name for p in folder.iterdir()}, {
            'manifest.json', 'sources.json', 'evidence.html', 'report.md', '人工指标核验表.xlsx'})
        wb = self.workbook()
        for name in ('实体核验', '溯源核验'):
            self.assertEqual(wb[name].max_row, 3)
            for row in wb[name].iter_rows(min_row=2):
                self.assertIsNone(row[7].value)
                self.assertEqual(row[9].value, '待核验')
                self.assertIn('保留待核验', row[21].value)
        wb.close()

    def test_invalid_rows_are_cleaned_before_workbook_and_manifest_export(self):
        from gpt_researcher.evaluation.manual_extraction import extract_published_rows
        rows = extract_published_rows(self.report, self.record['citation_map'],
                                      self.record['original_source_catalog'])
        rows.extend([{**rows[0], 'id': 'E90', 'claim': ' \u200b'},
                     {**rows[0], 'id': 'E91', 'name': '\ufffd' * 15}])
        with patch('gpt_researcher.evaluation.manual_review.extract_published_rows', return_value=rows):
            self.package = create_review_package(self.record, self.root)
        folder = Path(self.package['workbook']).parent
        manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(len(manifest['rows']), 4)
        self.assertEqual(manifest['export_cleanup']['removed_rows'], 2)
        self.assertEqual(manifest['extraction_audit'], [])
        wb = self.workbook()
        self.assertEqual(wb['实体核验'].max_row, 3)
        self.assertEqual(wb['提取审计'].max_row, 1)
        wb.close()

    def test_missing_source_rows_are_deleted_even_with_readable_evidence(self):
        self.report = self.report.replace('## 1 分析', 'LEAP维修50台。\n\n## 1 分析')
        self.record['published_markdown'] = self.report
        (self.root / 'outputs/测试.md').write_text(self.report, encoding='utf-8')
        self.package = create_review_package(self.record, self.root)
        folder = Path(self.package['workbook']).parent
        manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(manifest['export_cleanup']['missing_source_rows'], 1)
        self.assertTrue(all(row['source'].strip() for row in manifest['rows']))
        self.assertNotIn('LEAP', {row['name'] for row in manifest['rows']})
        self.assertEqual(manifest['extraction_audit'], [])
        self.assertIn('LEAP维修50台', (folder / 'report.md').read_text(encoding='utf-8'))
        wb = self.workbook()
        self.assertEqual(wb['实体核验'].max_row, 3)
        self.assertTrue(all(row[5] for row in list(wb['实体核验'].values)[1:]))
        wb.close()
        result = self.submit(self.fill(self.workbook()))
        self.assertEqual(result['overall'], '待完成')
        self.assertIsNone(result['entity']['requirement_met'])

    def test_workbook_retains_published_refs_and_has_no_candidate_sources(self):
        self.report += '\n[2] LEAP-1A未使用 https://example.test/unused\n'
        self.record['published_markdown'] = self.report
        self.record['citation_map']['[2]'] = {'url': 'https://example.test/unused'}
        self.record['original_source_catalog']['sources'].append({
            'locator': 'https://example.test/unused', 'pages': [{'text': 'LEAP-1A维修999台。'}]})
        (self.root / 'outputs/测试.md').write_text(self.report, encoding='utf-8')
        self.package = create_review_package(self.record, self.root)
        wb = self.workbook()
        self.addCleanup(wb.close)
        entities = list(wb['实体核验'].values)
        sources = list(wb['溯源核验'].values)
        self.assertEqual({row[2] for row in entities[1:]}, {'GTF', 'PW1100G'})
        self.assertTrue(all(row[4] == '[1]' for row in sources[1:]))
        self.assertTrue(all(row[2] == row[5] == 'https://example.test/a' for row in sources[1:]))
        self.assertTrue(all(row[6] for row in entities[1:] + sources[1:]))
        catalog = json.loads((Path(self.package['workbook']).parent / 'sources.json').read_text(encoding='utf-8'))
        self.assertEqual(len(catalog['sources']), 1)
        self.assertNotIn('example.test/unused', Path(self.package['materials']).read_text(encoding='utf-8').split('<h2>正文实际引用来源')[1])

    def test_legacy_workbook_still_imports_without_relabeling_its_policy(self):
        from gpt_researcher.evaluation.manual_workbook import LEGACY_HEADERS, fixed_values
        folder = Path(self.package['workbook']).parent
        manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
        manifest['policy_version'] = 'human-acceptance-v1'
        for row in manifest['rows']:
            row.pop('body_location', None)
        (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
        wb = self.fill(self.workbook())
        del wb['提取审计']
        for title, kind in [('实体核验', 'entity'), ('溯源核验', 'source'), ('漏项与争议', None)]:
            ws = wb[title]
            ws.delete_cols(14, 10)
            for column, header in enumerate(LEGACY_HEADERS, 1):
                ws.cell(1, column, header)
            originals = [row for row in manifest['rows'] if row['kind'] == kind]
            for index, row in enumerate(originals, 2):
                for column, value in enumerate(fixed_values(row), 1):
                    ws.cell(index, column, value)
        result = self.submit(wb)
        self.assertEqual(result['policy_version'], 'human-acceptance-v1')
        self.assertEqual(result['overall'], '达标')

    def test_v2_workbook_import_is_preserved(self):
        from gpt_researcher.evaluation.manual_workbook import HEADERS
        folder = Path(self.package['workbook']).parent
        manifest = json.loads((folder / 'manifest.json').read_text(encoding='utf-8'))
        manifest['policy_version'] = 'human-acceptance-v2-published-body'
        (folder / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False), encoding='utf-8')
        wb = self.fill(self.workbook())
        del wb['提取审计']
        for title in ('实体核验', '溯源核验', '漏项与争议'):
            wb[title].delete_cols(14, 10)
            self.assertEqual([c.value for c in wb[title][1]], HEADERS)
        self.assertEqual(self.submit(wb)['policy_version'], 'human-acceptance-v2-published-body')

    def test_structured_columns_and_audit_cannot_be_silently_changed(self):
        wb = self.fill(self.workbook())
        self.assertEqual(wb['实体核验']['N2'].value, '维修')
        self.assertEqual(wb['实体核验']['O2'].value, '120')
        self.assertEqual(wb['实体核验']['P2'].value, '台')
        self.assertTrue(wb['实体核验']['N2'].protection.locked)
        self.assertFalse(wb['实体核验']['J2'].protection.locked)
        wb['实体核验']['O2'] = '999'
        with self.assertRaisesRegex(ManualReviewError, '结构化'):
            self.submit(wb)
        wb = self.fill(self.workbook())
        del wb['提取审计']
        with self.assertRaisesRegex(ManualReviewError, '审计'):
            self.submit(wb)

    def test_all_sheets_allow_selection_and_copy_without_protection(self):
        wb = self.workbook()
        self.addCleanup(wb.close)
        self.assertTrue(all(not ws.protection.sheet for ws in wb))
        self.assertEqual(wb['实体核验']['F2'].value, 'https://example.test/a')
        self.assertEqual(wb['实体核验']['J2'].value, '待核验')
        self.assertTrue(wb['实体核验']['I2'].hyperlink.target.startswith('evidence.html#'))
        self.assertTrue(wb['指标汇总']['B2'].value.startswith('='))

    def test_audit_confirmation_required_before_completion(self):
        self.report = self.report.replace('## 1 分析', '建议提前核对LEAP。\n\n## 1 分析')
        self.record['published_markdown'] = self.report
        (self.root / 'outputs/测试.md').write_text(self.report, encoding='utf-8')
        self.package = create_review_package(self.record, self.root)
        result = self.submit(self.fill(self.workbook()))
        self.assertEqual(result['overall'], '待完成')
        wb = self.fill(self.workbook())
        for row in wb['任务与规则']:
            if row[0].value == '提取复核确认':
                row[1].value = '确认'
        self.assertEqual(self.submit(wb)['overall'], '达标')

    def test_completed_round_and_append_only_history(self):
        result = self.submit(self.fill(self.workbook()))
        self.assertEqual(result['overall'], '达标')
        self.assertEqual(result['entity']['accuracy'], 1)
        self.assertEqual(result['source']['accuracy'], 1)
        self.assertTrue(Path(result['paths']['word']).is_file())
        store = EvaluationRecordStore(self.root / 'records.json')
        store.append_run({**self.record, 'independent_judge': {'status': 'historical'}})
        store.append_manual_review(self.record['run_id'], result)
        store.append_manual_review(self.record['run_id'], result)
        saved = store.get_run(self.record['run_id'])
        self.assertEqual(len(saved['manual_reviews']), 2)
        self.assertEqual(saved['independent_judge']['status'], 'historical')
        self.assertEqual(len(store.list_runs()), 1)

    def test_pending_is_neither_zero_nor_pass(self):
        result = self.submit(self.fill(self.workbook(), '待核验'))
        self.assertIsNone(result['entity']['accuracy'])
        self.assertEqual(result['entity']['incorrect'], 0)
        self.assertEqual(result['overall'], '待完成')

    def test_missing_coverage_does_not_certify_subset(self):
        result = self.submit(self.fill(self.workbook(), confirm=False))
        self.assertIsNone(result['source']['requirement_met'])

    def test_zeros_and_threshold_boundaries(self):
        self.assertIsNone(score_rows([], True)['source']['accuracy'])
        rows = [{'kind': 'source', 'decision': '正确'}] * 49 + [{'kind': 'source', 'decision': '错误'}]
        self.assertTrue(score_rows(rows, True)['source']['requirement_met'])
        rows[0] = {'kind': 'source', 'decision': '错误'}
        self.assertFalse(score_rows(rows, True)['source']['requirement_met'])

    def test_changed_word_file_invalidates_review(self):
        word = self.root / 'outputs/report.docx'
        word.write_bytes(b'original')
        self.record['word_path'] = 'outputs/report.docx'
        self.package = create_review_package(self.record, self.root)
        word.write_bytes(b'modified')
        with self.assertRaisesRegex(ManualReviewError, '版本'):
            self.submit(self.fill(self.workbook()))

    def test_removed_rows_and_modified_fixed_cells_rejected(self):
        for mode in ['delete', 'edit', 'duplicate', 'formula']:
            with self.subTest(mode=mode):
                wb = self.fill(self.workbook())
                ws = wb['实体核验']
                if mode == 'delete':
                    ws.delete_rows(2)
                elif mode == 'edit':
                    ws['D2'] = '伪造原句'
                elif mode == 'duplicate':
                    ws.append([c.value for c in ws[2]])
                else:
                    ws['J2'] = '=1+1'
                with self.assertRaises(ManualReviewError):
                    self.submit(wb)

    def test_extra_entity_and_unknown_verdict(self):
        wb = self.fill(self.workbook())
        ws = wb['漏项与争议']
        ws.append(['A1', 'entity:型号', '另一核验属性', 'GTF维修120台[1]。', '[1]', 'https://example.test/a', '', '', '',
                   '正确', '人工补充', 'GTF维修120台。', '网页第2段'])
        result = self.submit(wb)
        self.assertEqual(result['entity']['total'], self.package['entity_count'] + 1)
        wb = self.fill(self.workbook())
        wb['实体核验']['J2'] = '基本正确'
        with self.assertRaises(ManualReviewError):
            self.submit(wb)

    def test_manual_addition_resolves_published_markdown_citation_links(self):
        self.report = self.report.replace('[1]。', '[[1]](#ref-1)。')
        self.record['published_markdown'] = self.report
        (self.root / 'outputs/测试.md').write_text(self.report, encoding='utf-8')
        self.package = create_review_package(self.record, self.root)
        wb = self.fill(self.workbook())
        wb['漏项与争议'].append(['A1', 'entity:参数', '补充属性', 'GTF维修120台。', '', '', '', '', '',
                                  '正确', '补充候选', 'GTF维修120台。', '原文第2页'])
        self.assertEqual(self.submit(wb)['overall'], '达标')

    def test_exclusion_requires_distinct_reviewer(self):
        wb = self.fill(self.workbook())
        wb['实体核验']['J2'] = '排除'
        result = self.submit(wb)
        self.assertEqual(result['overall'], '待完成')
        wb = self.fill(self.workbook())
        wb['实体核验']['J2'] = '排除'
        for row in wb['任务与规则'].iter_rows(min_row=2):
            if row[0].value == '复核人':
                row[1].value = '另一人'
            elif row[0].value == '排除项复核确认':
                row[1].value = '确认'
        result = self.submit(wb)
        self.assertEqual(result['overall'], '达标')

    def test_untrusted_text_not_formula_or_html(self):
        self.record['task'] = '=HYPERLINK("https://bad.test")'
        self.record['original_source_catalog']['sources'][0]['pages'][0]['text'] += '<script>alert(1)</script>'
        package = create_review_package(self.record, self.root)
        wb = load_workbook(package['workbook'])
        self.addCleanup(wb.close)
        task = next(row[1] for row in wb['任务与规则'] if row[0].value == '任务')
        self.assertEqual(task.data_type, 's')
        self.assertNotIn('<script>', Path(package['materials']).read_text(encoding='utf-8'))

    def test_ten_distinct_rounds_all_must_pass(self):
        result = self.submit(self.fill(self.workbook()))
        records = []
        for index in range(10):
            record = copy.deepcopy(self.record)
            record['run_id'] = f'run-{index}'
            record['manual_reviews'] = [{**result, 'run_id': record['run_id']}]
            records.append(record)
        self.assertTrue(summarize_runs(records, self.root)['requirement_met'])
        self.assertFalse(summarize_runs(records[:9], self.root)['requirement_met'])
        records[-1]['manual_reviews'][-1] = {**result, 'overall': '未达标'}
        self.assertFalse(summarize_runs(records, self.root)['requirement_met'])
        with self.assertRaises(ManualReviewError):
            summarize_runs([records[0], records[0]], self.root)

    def test_material_tampering_is_rejected(self):
        Path(self.package['materials']).write_text('changed', encoding='utf-8')
        with self.assertRaisesRegex(ManualReviewError, '证据快照'):
            self.submit(self.fill(self.workbook()))

    def test_wrong_run_and_missing_files_rejected(self):
        with self.assertRaises(ManualReviewError):
            self.submit(self.fill(self.workbook()), {**self.record, 'run_id': 'another'})
        (self.root / 'outputs/测试.md').unlink()
        with self.assertRaises(ManualReviewError):
            self.submit(self.fill(self.workbook()))

    def test_actual_word_export(self):
        from docx import Document
        wb = self.fill(self.workbook())
        stream = BytesIO()
        wb.save(stream)
        wb.close()
        result = import_review(self.record, stream.getvalue(), self.root)
        self.assertTrue(result['paths']['word'], result.get('export_error'))
        document = Document(result['paths']['word'])
        self.assertIn('人工指标测评报告', '\n'.join(p.text for p in document.paragraphs))
        self.assertTrue(document.tables)

    def test_concurrent_cli_and_server_updates_do_not_lose_history(self):
        store = EvaluationRecordStore(self.root / 'records.json')
        store.append_run(self.record)
        code = (
            "import sys; sys.path.insert(0,sys.argv[1]); "
            "from gpt_researcher.evaluation.records import EvaluationRecordStore; "
            "s=EvaluationRecordStore(sys.argv[2]); "
            "[s.append_reevaluation('test-001', {'worker':sys.argv[3], 'iteration':i}) for i in range(10)]"
        )
        processes = [subprocess.Popen([sys.executable, '-c', code, str(Path(__file__).resolve().parents[1]),
                                       str(store.path), str(i)], stdout=subprocess.PIPE, stderr=subprocess.PIPE)
                     for i in range(3)]
        for process in processes:
            _, error = process.communicate(timeout=60)
            self.assertEqual(process.returncode, 0, error.decode(errors='replace'))
        self.assertEqual(len(store.get_run('test-001')['reevaluations']), 30)


if __name__ == '__main__':
    unittest.main()
